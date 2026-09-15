"""Visible, mandatory identification for all generated demonstration documents."""
import re

SAMPLE_LABEL = 'SAMPLE — NOT A BANK-ISSUED DOCUMENT'


def sample_html(html: str) -> str:
    html = re.sub(r'<div data-sample-document="true"[^>]*>SAMPLE [^<]*</div>', '', html)
    banner = '<div data-sample-document="true" style="color:#a40000;font:bold 12pt Arial;text-align:center;padding:6pt;border:2px solid #a40000">' + SAMPLE_LABEL + '</div>'
    return re.sub(r'(<body\b[^>]*>)', lambda m: m.group(1) + banner, html, count=1, flags=re.I) if re.search(r'<body\b', html, re.I) else banner + html


def mark_workbook(workbook) -> None:
    from openpyxl.styles import Font
    for sheet in workbook.worksheets:
        sheet.cell(1, 1, SAMPLE_LABEL + '\n' + str(sheet.cell(1, 1).value or '')).font = Font(bold=True, color='A40000', size=12)
        from copy import copy
        alignment = copy(sheet.cell(1, 1).alignment)
        alignment.wrap_text = True
        sheet.cell(1, 1).alignment = alignment
        sheet.row_dimensions[1].height = max(40, sheet.row_dimensions[1].height or 0)
        sheet.oddHeader.center.text = '&B SAMPLE - NOT A BANK-ISSUED DOCUMENT'
        sheet.evenHeader.center.text = sheet.oddHeader.center.text
        sheet.firstHeader.center.text = sheet.oddHeader.center.text
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_area = sheet.calculate_dimension()


def mark_document(document) -> None:
    from docx.shared import Pt, RGBColor
    paragraph = document.paragraphs[0].insert_paragraph_before(SAMPLE_LABEL) if document.paragraphs else document.add_paragraph(SAMPLE_LABEL)
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            label = header.add_paragraph(SAMPLE_LABEL)
            label.alignment = 1
            for run in label.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string('A40000')
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string('A40000')
