"""Office layout, persistent editing, legacy conversion and sample export tests."""
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import app
from .document_layout import word_html
from .letterheads import load_letterhead, save_letterhead
from .sample_documents import mark_document, mark_workbook, sample_html


class OfficeFormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='sample-office-tests-')
        self.root = Path(self.temp.name)
        self.patch = patch.object(app, 'CUSTOM_TEMPLATE_ROOT', self.root / 'custom')
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def test_word_order_spacing_runs_and_sample_mark(self):
        from docx import Document
        from docx.shared import Pt
        document = Document()
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        p.add_run('Before  table').bold = True
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Name'
        table.cell(0, 1).text = '{{customer_name}}'
        document.add_paragraph('After\ttable')
        path = self.root / 'sample.docx'
        document.save(path)
        html = word_html(path, editable_keys=True)
        self.assertLess(html.index('Before  table'), html.index('<table'))
        self.assertLess(html.index('</table>'), html.index('After'))
        self.assertIn('margin-left:24pt', html)
        self.assertIn('font-weight:bold', html)
        self.assertIn('data-template-key="p:0"', html)
        mark_document(document)
        self.assertIn('SAMPLE', document.paragraphs[0].text)
        self.assertTrue(all('SAMPLE' in s.header.paragraphs[-1].text for s in document.sections))
        self.assertEqual(sample_html(sample_html(html)), sample_html(html))

    def test_save_all_is_atomic_and_does_not_mutate_bundled_workbook(self):
        from openpyxl import Workbook, load_workbook
        workbook = Workbook()
        worksheet = workbook.active
        worksheet['A1'] = 'Old A'
        worksheet['B1'] = 'Old B'
        worksheet['C2'] = '=1+2'
        worksheet.column_dimensions['B'].width = 32
        worksheet.row_dimensions[1].height = 29
        bundled = self.root / 'bundled'
        bundled.mkdir()
        path = bundled / 'Model Sample.xlsx'
        workbook.save(path)
        original = path.read_bytes()
        detail = app.template_detail('statement', 'Model Sample', str(bundled), 777)
        self.assertTrue(detail['editable'])
        key = worksheet.title + '!'
        edits = [{'key': key+'A1','text':'New A'}, {'key':key+'B1','text':'New B'}]
        app.update_template_batch('statement', 'Model Sample', str(bundled), 777, edits)
        custom = app.resolve_template_entry('statement', 'Model Sample', str(bundled), 777).path
        saved = load_workbook(custom)
        self.assertEqual(saved.active['A1'].value, 'New A')
        self.assertEqual(saved.active['B1'].value, 'New B')
        self.assertEqual(saved.active['C2'].value, '=1+2')
        self.assertEqual(saved.active.column_dimensions['B'].width, 32)
        self.assertEqual(saved.active.row_dimensions[1].height, 29)
        before = custom.read_bytes()
        with self.assertRaises(ValueError):
            app.update_template_batch('statement', 'Model Sample', str(bundled), 777, [{'key':key+'A1','text':'Must roll back'}, {'key':'invalid','text':'bad'}])
        self.assertEqual(custom.read_bytes(), before)
        self.assertEqual(path.read_bytes(), original)

    def test_word_cell_edit_does_not_duplicate_old_paragraphs(self):
        from docx import Document
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0,0)
        cell.text = 'First'
        cell.add_paragraph('Second')
        path = self.root / 'cell.docx'
        document.save(path)
        app.update_word_template_item(path, 't:0:r:0:c:0', 'Changed\nSecond', {})
        self.assertEqual(Document(path).tables[0].cell(0,0).text, 'Changed\nSecond')

    def test_letterheads_are_scoped_and_invalid_margins_rejected(self):
        image='data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='
        value = save_letterhead(self.root, 1, 'statement', 'Model 1', {'image':image})
        self.assertEqual(load_letterhead(self.root,1,'statement','Model 1'),value)
        for user,kind,name in [(2,'statement','Model 1'),(1,'certificate','Model 1'),(1,'statement','Model 2')]:
            self.assertEqual(load_letterhead(self.root,user,kind,name),{})
        for bad in [float('nan'), -1, 91]:
            with self.assertRaises(ValueError):
                save_letterhead(self.root,1,'statement','Model 1',{'top':bad})
        self.assertEqual(save_letterhead(self.root,1,'statement','Model 1',{'remove':True}),{})

    def test_workbook_sample_mark_preserves_title_merges_and_formulas(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.merge_cells('A1:D1')
        sheet['A1']='Example Statement'
        sheet['D4']='=SUM(D2:D3)'
        mark_workbook(workbook)
        self.assertIn('SAMPLE',sheet['A1'].value)
        self.assertIn('A1:D1',[str(x) for x in sheet.merged_cells.ranges])
        self.assertEqual(sheet['D4'].value,'=SUM(D2:D3)')
        self.assertIn('SAMPLE',sheet.oddHeader.center.text)

    @unittest.skipUnless(os.name == 'nt' and os.environ.get('SGV2_TEST_OFFICE') == '1', 'Set SGV2_TEST_OFFICE=1 with Microsoft Office installed')
    def test_legacy_excel_and_word_conversion(self):
        from openpyxl import Workbook, load_workbook
        from docx import Document
        from .office_formats import convert_office, editable_office_copy
        workbook = Workbook()
        workbook.active['A1'] = 'SAMPLE legacy worksheet'
        workbook.save(self.root/'sample.xlsx')
        convert_office(self.root/'sample.xlsx', self.root/'sample.xls')
        converted = editable_office_copy(self.root/'sample.xls')
        self.assertEqual(load_workbook(converted).active['A1'].value, 'SAMPLE legacy worksheet')
        document=Document()
        document.add_paragraph('SAMPLE legacy certificate')
        document.save(self.root/'sample.docx')
        convert_office(self.root/'sample.docx',self.root/'sample.doc')
        self.assertEqual(Document(editable_office_copy(self.root/'sample.doc')).paragraphs[0].text,'SAMPLE legacy certificate')

    @unittest.skipUnless(os.name == 'nt' and os.environ.get('SGV2_TEST_OFFICE') == '1', 'Set SGV2_TEST_OFFICE=1 with Microsoft Office installed')
    def test_native_exports_resize_rows_and_populate_certificate(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import PatternFill
        from docx import Document
        from .exporters import export_statement, export_certificate
        form = app.default_form_values()
        form.update(customer_name='SAMPLE Customer', customer_address='SAMPLE Address', reference_no='SAMPLE-REF-77')
        _, _, _, payload = app.build_export_input({'config':form,'generated_seed':'123456','rate_mode':'Manual','manual_rate':'145.25'})
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Date','Description','Debit','Credit','Balance'])
        for _ in range(4):
            sheet.append(['2026-01-05','Old sample row',0,100,100])
            for cell in sheet[sheet.max_row]: cell.fill = PatternFill('solid',fgColor='FFEECC')
        sheet.append(['','Total',0,0,0])
        sheet.append(['','Prepared by SAMPLE reviewer'])
        source = self.root/'rows.xlsx'
        workbook.save(source)
        original = source.read_bytes()
        for count in (8,2):
            case = dict(payload)
            case['statement_rows'] = payload['statement_rows'][:count]
            output = self.root/f'rows-{count}.xlsx'
            export_statement(source,output,case)
            result = load_workbook(output)
            rows = list(result.active.values)
            total = next(i for i,row in enumerate(rows) if len(row)>1 and row[1]=='Total')
            self.assertEqual(total, count+2)  # Mandatory sample row plus column header.
            self.assertIn('Prepared by SAMPLE reviewer',str(rows[total+1]))
            self.assertIn('SAMPLE',str(rows[0]))
            self.assertEqual(result.active.cell(3,2).fill.fgColor.rgb[-6:],'FFEECC')
        self.assertEqual(source.read_bytes(),original)
        document = Document()
        p = document.add_paragraph()
        p.add_run('Name: ').bold=True
        p.add_run('{{customer_name}}').italic=True
        table=document.add_table(rows=0,cols=2)
        for label in ['Address','Reference No.','Issue Date','Total Balance','Exchange Rate on Issue Date','In Words NPR']:
            row=table.add_row();row.cells[0].text=label;row.cells[1].text='OLD SAMPLE VALUE'
        word_source=self.root/'certificate.docx';document.save(word_source)
        output=self.root/'certificate-result.docx'
        export_certificate(word_source,output,payload)
        result=Document(output)
        content='\n'.join(p.text for p in result.paragraphs)+'\n'+'\n'.join(c.text for r in result.tables[0].rows for c in r.cells)
        for value in ['SAMPLE Customer','SAMPLE Address','SAMPLE-REF-77','145.25',payload['certificate']['total_balance_npr_text'],payload['certificate']['balance_words_npr']]:
            self.assertIn(value,content)
        self.assertNotIn('OLD SAMPLE VALUE',content)
        self.assertTrue(result.paragraphs[1].runs[0].bold)
        self.assertIn('SAMPLE',result.sections[0].header.paragraphs[-1].text)


if __name__ == '__main__':
    unittest.main()
