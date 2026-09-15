from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from uuid import uuid4

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:  # pragma: no cover - depends on local Python install
    Workbook = None
    Alignment = Border = Font = PatternFill = Side = None

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:  # pragma: no cover - depends on local Python install
    Document = None
    WD_ALIGN_PARAGRAPH = Pt = None

from .exchange_rate import ExchangeRateLookupError, ExchangeRateResult, fetch_usd_npr_rate
from .generator import StatementConfig, StatementResult, validate_transaction_dates
from .utils import (
    amount_to_words_npr,
    amount_to_words_usd,
    format_amount,
    format_long_date,
    format_slash_date,
    iso_date,
    safe_filename,
    write_json,
)


STATEMENT_EXTENSIONS = {".xlsx", ".xls"}
CERTIFICATE_EXTENSIONS = {".docx", ".doc"}


@dataclass(slots=True)
class TemplateEntry:
    name: str
    path: Path
    source: str = "bundled"
    editable: bool = False


@dataclass(slots=True)
class TemplateCatalog:
    statement_templates: list[TemplateEntry]
    certificate_templates: list[TemplateEntry]


def scan_template_directory(directory: Path) -> TemplateCatalog:
    statement_templates: list[TemplateEntry] = []
    certificate_templates: list[TemplateEntry] = []
    if not directory.exists():
        return TemplateCatalog(statement_templates, certificate_templates)

    for path in sorted(directory.iterdir(), key=lambda item: item.name.lower()):
        suffix = path.suffix.lower()
        if suffix in STATEMENT_EXTENSIONS:
            statement_templates.append(TemplateEntry(path.stem, path))
        elif suffix in CERTIFICATE_EXTENSIONS:
            certificate_templates.append(TemplateEntry(path.stem, path))
    return TemplateCatalog(statement_templates, certificate_templates)


def resolve_exchange_rate(
    issue_date: date,
    mode: str,
    manual_rate: float | None,
    rate_type: str,
    timeout: int = 20,
) -> ExchangeRateResult:
    if mode == "manual":
        if not manual_rate or manual_rate <= 0:
            raise ExchangeRateLookupError("Enter a valid manual USD/NPR exchange rate before exporting.")
        return ExchangeRateResult(
            rate=manual_rate,
            rate_type=rate_type.lower(),
            source_date=issue_date,
            source_label="Manual USD/NPR Rate",
        )
    return fetch_usd_npr_rate(issue_date, rate_type=rate_type, timeout=timeout)


def build_payload(
    config: StatementConfig,
    result: StatementResult,
    exchange_rate: ExchangeRateResult,
) -> dict:
    errors = validate_transaction_dates(config, [asdict(row) for row in result.rows])
    if errors:
        raise ValueError(str(errors[0]["message"]) + " Recalculate the statement before exporting.")

    def num_text(value: float, decimals: int = 2) -> str:
        return f"{float(value):.{decimals}f}"

    total_balance_npr = result.final_balance
    equivalent_usd = round(total_balance_npr / exchange_rate.rate, 2)
    total_debit = round(sum(float(row.debit) for row in result.rows), 2)
    total_credit = round(sum(float(row.credit) for row in result.rows), 2)
    rows = [
        {
            "date": iso_date(row.date),
            "txn_date": iso_date(row.date),
            "value_date": iso_date(row.date),
            "description": row.description,
            "cheque_no": row.cheque_no,
            "debit": num_text(row.debit),
            "credit": num_text(row.credit),
            "balance": num_text(row.balance),
            "debit_text": format_amount(row.debit) if row.debit else "",
            "credit_text": format_amount(row.credit) if row.credit else "",
            "balance_text": format_amount(row.balance),
            "category": row.category,
            "is_system": row.is_system,
        }
        for row in result.rows
    ]
    return {
        "account": {
            "bank_name": config.bank_name,
            "branch_name": config.branch_name,
            "customer_name": config.customer_name,
            "customer_address": config.customer_address,
            "account_number": config.account_number,
            "account_type": config.account_type,
            "member_id": config.member_id,
            "currency": config.currency,
            "reference_no": config.reference_no,
            "opening_date_iso": iso_date(config.opening_date) if config.opening_date else "",
            "opening_date_slash": format_slash_date(config.opening_date) if config.opening_date else "",
        },
        "statement": {
            "period_from_iso": iso_date(config.start_date),
            "period_to_iso": iso_date(result.last_transaction_date),
            "period_from_slash": format_slash_date(config.start_date),
            "period_to_slash": format_slash_date(result.last_transaction_date),
            "period_label_slash": f"{format_slash_date(config.start_date)} to {format_slash_date(result.last_transaction_date)}",
            "period_label_iso": f"{iso_date(config.start_date)} to {iso_date(result.last_transaction_date)}",
            "issue_date_iso": iso_date(result.issue_date),
            "issue_date_slash": format_slash_date(result.issue_date),
            "issue_date_long": format_long_date(result.issue_date, ordinal=False),
            "issue_date_ordinal": format_long_date(result.issue_date, ordinal=True),
            "as_of_iso": iso_date(result.last_transaction_date),
            "as_of_slash": format_slash_date(result.last_transaction_date),
            "as_of_long": format_long_date(result.last_transaction_date, ordinal=False),
            "as_of_ordinal": format_long_date(result.last_transaction_date, ordinal=True),
            "opening_business_date_iso": iso_date(result.opening_business_date),
            "ending_business_date_iso": iso_date(result.ending_business_date),
            "include_cheque_column": config.include_cheque_column,
            "date_column_mode": config.date_column_mode if config.date_column_mode in {"single", "txn_value"} else "single",
        },
        "rates": {
            "interest_rate": num_text(config.interest_rate, 4).rstrip("0").rstrip("."),
            "tax_rate": num_text(config.tax_rate, 4).rstrip("0").rstrip("."),
            "usd_npr": num_text(exchange_rate.rate, 4),
            "usd_npr_text": f"{exchange_rate.rate:,.2f}",
            "rate_type": exchange_rate.rate_type,
            "source_label": exchange_rate.source_label,
            "source_date_iso": iso_date(exchange_rate.source_date),
        },
        "summary": {
            "total_deposits": num_text(result.summary.total_deposits),
            "total_withdrawals": num_text(result.summary.total_withdrawals),
            "total_debit": num_text(total_debit),
            "total_credit": num_text(total_credit),
            "total_debit_text": format_amount(total_debit),
            "total_credit_text": format_amount(total_credit),
            "total_interest": num_text(result.summary.total_interest),
            "total_tax": num_text(result.summary.total_tax),
            "deposit_count": str(result.summary.deposit_count),
            "withdrawal_count": str(result.summary.withdrawal_count),
            "row_count": str(len(result.rows)),
            "final_balance": num_text(total_balance_npr),
            "final_balance_text": format_amount(total_balance_npr),
        },
        "certificate": {
            "total_balance_npr": num_text(total_balance_npr),
            "total_balance_npr_text": format_amount(total_balance_npr),
            "equivalent_usd": num_text(equivalent_usd),
            "equivalent_usd_text": format_amount(equivalent_usd),
            "balance_words_npr": amount_to_words_npr(total_balance_npr),
            "balance_words_usd": amount_to_words_usd(equivalent_usd),
            "authorization_details": "Authorized Signature",
        },
        "statement_rows": rows,
    }


def default_output_name(kind: str, customer_name: str, template_name: str, issue_date: date, suffix: str) -> str:
    name = safe_filename(customer_name)
    template = safe_filename(template_name)
    return f"{kind}_{name}_{issue_date.isoformat()}_{template}{suffix}"


def _resource_root() -> Path:
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent


def _resource_file(name: str) -> Path:
    root = _resource_root()
    candidates = [
        root / name,
        root / "statement_generator" / name,
        Path(__file__).resolve().parent / name,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def _run_export(mode: str, template_path: Path, output_path: Path, payload: dict) -> None:
    script_path = _resource_file("office_export.ps1")
    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    temp_root = _resource_root() / "_web_temp"
    temp_root.mkdir(parents=True, exist_ok=True)
    temp_dir = temp_root / f"statement_generator_{uuid4().hex}"
    temp_dir.mkdir(parents=True, exist_ok=True)
    try:
        payload_path = temp_dir / "payload.json"
        write_json(payload_path, payload)
        command = [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-Mode",
            mode,
            "-TemplatePath",
            str(template_path),
            "-OutputPath",
            str(output_path),
            "-PayloadPath",
            str(payload_path),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
            creationflags=creationflags,
        )
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
    if completed.returncode != 0:
        message = completed.stderr.strip() or completed.stdout.strip() or "Office export failed."
        raise RuntimeError(message)


def export_statement(template_path: Path, output_path: Path, payload: dict) -> None:
    _run_export("statement", template_path, output_path, payload)


def export_certificate(template_path: Path, output_path: Path, payload: dict) -> None:
    _run_export("certificate", template_path, output_path, payload)


def export_normal_statement(output_path: Path, payload: dict) -> None:
    if Workbook is None:
        raise RuntimeError("Normal Excel export dependency is missing in this environment.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"

    title_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="9AA5B1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    statement_payload = payload.get("statement", {})
    include_cheque = bool(statement_payload.get("include_cheque_column", True))
    two_date_columns = statement_payload.get("date_column_mode") == "txn_value"
    total_columns = (2 if two_date_columns else 1) + 4 + (1 if include_cheque else 0)
    merge_end_column = chr(ord("A") + total_columns - 1)
    ws.merge_cells(f"A1:{merge_end_column}1")
    ws["A1"] = "BANK STATEMENT"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A3"] = "Name"
    ws["B3"] = payload["account"]["customer_name"]
    ws["D3"] = "Account No."
    ws["E3"] = payload["account"]["account_number"]
    ws["A4"] = "Address"
    ws["B4"] = payload["account"]["customer_address"]
    ws["D4"] = "Account Type"
    ws["E4"] = payload["account"]["account_type"]
    ws["A5"] = "Statement Period"
    ws["B5"] = payload["statement"]["period_label_slash"]
    ws["D5"] = "Issue Date"
    ws["E5"] = payload["statement"]["issue_date_slash"]
    ws["A6"] = "Interest Rate"
    ws["B6"] = f"{payload['rates']['interest_rate']}%"
    ws["D6"] = "Tax Rate"
    ws["E6"] = f"{payload['rates']['tax_rate']}%"

    headers = ["TXN Date", "Value Date"] if two_date_columns else ["Date"]
    headers.append("Description")
    if include_cheque:
        headers.append("Cheque No.")
    headers.extend(["Debit", "Credit", "Balance"])
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=8, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    row_index = 9
    for item in payload["statement_rows"]:
        date_cell = ws.cell(row=row_index, column=1, value=datetime.strptime(item.get("txn_date") or item["date"], "%Y-%m-%d").date())
        date_cell.number_format = "yyyy-mm-dd"
        date_cell.alignment = Alignment(wrap_text=False)
        description_column = 2
        if two_date_columns:
            value_date_cell = ws.cell(row=row_index, column=2, value=datetime.strptime(item.get("value_date") or item["date"], "%Y-%m-%d").date())
            value_date_cell.number_format = "yyyy-mm-dd"
            value_date_cell.alignment = Alignment(wrap_text=False)
            description_column = 3
        description_cell = ws.cell(row=row_index, column=description_column, value=item["description"])
        description_cell.number_format = "@"
        amount_start_column = description_column + 1
        if include_cheque:
            cheque_text = str(item["cheque_no"]).strip()
            cheque_value = int(cheque_text) if cheque_text.isdigit() else None
            cheque_cell = ws.cell(row=row_index, column=description_column + 1, value=cheque_value)
            cheque_cell.number_format = "0"
            amount_start_column = description_column + 2
        debit = float(item["debit"])
        credit = float(item["credit"])
        balance = float(item["balance"])
        ws.cell(row=row_index, column=amount_start_column, value=debit if debit > 0 else None)
        ws.cell(row=row_index, column=amount_start_column + 1, value=credit if credit > 0 else None)
        ws.cell(row=row_index, column=amount_start_column + 2, value=balance)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_index, column=col)
            cell.border = border
            if col >= amount_start_column and cell.value is not None:
                cell.number_format = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
        row_index += 1

    widths = {"A": 14, "B": 14 if two_date_columns else 42, "C": 42 if two_date_columns else (16 if include_cheque else 14), "D": 16 if include_cheque else 14, "E": 14, "F": 14, "G": 16}
    for column, width in widths.items():
        if ord(column) - ord("A") + 1 > total_columns:
            continue
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A9"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.print_area = f"A1:{merge_end_column}{max(8, row_index - 1)}"

    from .sample_documents import mark_workbook
    mark_workbook(wb)
    wb.save(output_path)


def export_normal_certificate(output_path: Path, payload: dict) -> None:
    if Document is None:
        raise RuntimeError("Normal Word export dependency is missing in this environment.")

    doc = Document()
    try:
        from docx.shared import Mm  # noqa: E402

        for section in doc.sections:
            section.page_width = Mm(210)
            section.page_height = Mm(297)
            section.left_margin = Mm(12.7)
            section.right_margin = Mm(12.7)
            section.top_margin = Mm(12.7)
            section.bottom_margin = Mm(12.7)
    except Exception:
        pass
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BALANCE CERTIFICATE")
    run.bold = True
    run.font.size = Pt(16)

    intro = doc.add_paragraph()
    intro.add_run("To whom it may concern\n").bold = True
    intro.add_run(
        f"This is to certify that the balance in the credit of the under mentioned account holder as on "
        f"{payload['statement']['as_of_ordinal']} is as follows."
    )

    lines = [
        f"Name: {payload['account']['customer_name']}",
        f"Address: {payload['account']['customer_address']}",
        f"Account Number: {payload['account']['account_number']}",
        f"Account Type: {payload['account']['account_type']}",
        f"Currency: {payload['account']['currency']}",
        f"Total Balance: NPR {payload['certificate']['total_balance_npr_text']}",
        f"Equivalent to USD: {payload['certificate']['equivalent_usd_text']}",
        f"In Words USD: {payload['certificate']['balance_words_usd']}",
        f"Exchange Rate on Issue Date: 1 USD = NPR {payload['rates']['usd_npr_text']}",
        f"Issue Date: {payload['statement']['issue_date_slash']}",
        f"Reference No.: {payload['account']['reference_no']}",
        f"In Words NPR: {payload['certificate']['balance_words_npr']}",
    ]
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.add_run(line)

    closing = doc.add_paragraph()
    closing.add_run(
        "This certificate has been issued at the request of the account holder without obligation on the part of the institution."
    )

    signature = doc.add_paragraph("\n\nAuthorized Signature")
    signature.alignment = WD_ALIGN_PARAGRAPH.LEFT

    from .sample_documents import mark_document
    mark_document(doc)
    doc.save(output_path)
