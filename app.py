from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
import urllib.request
import webbrowser
from datetime import date, datetime, timedelta
from email.parser import BytesParser
from email.policy import default as email_policy
from html import escape
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from uuid import uuid4

from auth_store import AuthStore, AuthUser


APP_ROOT = Path(__file__).resolve().parent
for vendor_path in (APP_ROOT / ".deps", APP_ROOT / ".vendor", APP_ROOT.parent / ".vendor"):
    if vendor_path.exists():
        sys.path.insert(0, str(vendor_path))
sys.path.insert(0, str(APP_ROOT))

from statement_generator.exchange_rate import (  # noqa: E402
    ExchangeRateLookupError,
    NRB_FOREX_DOCS,
    NRB_FOREX_PAGE,
)
from statement_generator.exporters import (  # noqa: E402
    CERTIFICATE_EXTENSIONS,
    STATEMENT_EXTENSIONS,
    TemplateCatalog,
    TemplateEntry,
    build_payload,
    default_output_name,
    export_certificate,
    export_statement,
    resolve_exchange_rate,
    scan_template_directory,
)
from statement_generator.generator import (  # noqa: E402
    StatementConfig,
    generate_statement,
    names_from_text,
    posting_date_for_period,
    recalculate_edited_statement,
    validate_edited_statement,
    validate_transaction_dates,
)
from statement_generator.rounding import parse_percentages
from statement_generator.office_formats import convert_office, editable_office_copy
from statement_generator.document_layout import word_html
from statement_generator.sample_documents import sample_html
from statement_generator.letterheads import load_letterhead, save_letterhead
from statement_generator.holidays import MANUAL_HOLIDAY_SEED_VERSION, SUNDAY_HOLIDAY_START, is_recurring_holiday, manual_holiday_dates  # noqa: E402
from statement_generator.importers import import_xlsx_statement  # noqa: E402
from statement_generator.selftest import run_tests  # noqa: E402
from statement_generator.utils import format_amount, json_default, next_business_day, parse_iso_date, resolve_business_day, round_money, safe_filename  # noqa: E402


DEFAULT_TEMPLATE_DIR = Path(r"D:\Finance Doc\Format")
HAMROPATRO_ENGLISH_CALENDAR_URL = "https://english.hamropatro.com/calendar/"
STATE_FILE = APP_ROOT / "web_statement_generator_state.json"
USER_STATE_ROOT = APP_ROOT / "user_state"
AUTH_DB_FILE = APP_ROOT / "web_statement_generator.db"
PERSISTENT_RULES_SCHEMA_VERSION = 6
QUARTER_MONTHS = (1, 4, 7, 10)
AUTO_REFRESH_SUCCESS_INTERVAL = timedelta(hours=6)
AUTO_REFRESH_FAILURE_BACKOFF = timedelta(minutes=30)
AUTO_HAMROPATRO_CALENDAR_TIMEOUT = 4
MANUAL_HAMROPATRO_TIMEOUT = 10
PERSISTENT_META_KEYS = ("last_auto_refresh_attempt_at", "last_auto_refresh_success_at")
BS_MONTH_LABELS = {
    1: "Baishakh",
    2: "Jestha",
    3: "Ashadh",
    4: "Shrawan",
    5: "Bhadra",
    6: "Ashwin",
    7: "Kartik",
    8: "Mangsir",
    9: "Paush",
    10: "Magh",
    11: "Falgun",
    12: "Chaitra",
}
DESCRIPTION_MODE_OPTIONS = {
    "Name + Label": "name_plus_label",
    "Label + Name": "label_plus_name",
    "Name Only": "name_only",
    "Label Only": "label_only",
    "Extra Text Only": "extra_only",
    "Label + Extra Text + Name": "label_extra_name",
    "Label + Name + Extra Text": "label_name_extra",
    "Name + Label + Extra Text": "name_label_extra",
    "Name + Extra Text + Label": "name_extra_label",
    "Extra Text + Label + Name": "extra_label_name",
    "Extra Text + Name + Label": "extra_name_label",
}
DESCRIPTION_MODE_LABELS = {value: key for key, value in DESCRIPTION_MODE_OPTIONS.items()}
AMOUNT_ROUNDING_OPTIONS = {
    "Automatic (70% / 20% / 10%)": "automatic",
    "Customized percentages": "custom",
    "Rounding Figure 5": "round_5",
    "Rounding Figure 10": "round_10",
    "Rounding Figure 50": "round_50",
    "Rounding Figure 100": "round_100",
    "Rounding Figure 500": "round_500",
    "Rounding Figure 1000": "round_1000",
    "Rounding Figure 1000 and 500": "round_1000_500",
    "Rounding Figure 1000, 500 and 100": "round_1000_500_100",
    "Rounding Figure 1000, 500, 100 and 50": "round_1000_500_100_50",
    "Rounding Figure 1000, 500, 100, 50 and 10": "round_1000_500_100_50_10",
    "Rounding Figure 1000, 500, 100, 50, 10 and 5": "round_1000_500_100_50_10_5",
}
DEFAULT_DEPOSIT_NAMES = "Self\nKaruna\nKrishna\nManisha"
DEFAULT_WITHDRAWAL_NAMES = "Self\nKabita Thapa\nKamala Pandey"
STATE_LOCK = threading.Lock()
LOGIN_ATTEMPT_FILE = APP_ROOT / "login_attempts.json"
LOGIN_ATTEMPT_LOCK = threading.Lock()
LOGIN_MAX_FAILURES = 6
LOGIN_WINDOW_SECONDS = 900
WEB_TEMP_ROOT = APP_ROOT / "_web_runtime_temp"
CUSTOM_TEMPLATE_ROOT = APP_ROOT / "custom_templates"
HIDDEN_TEMPLATE_FILE = "_hidden_templates.json"
TEMPLATE_META_SUFFIX = ".meta.json"
AUTH_STORE: AuthStore | None = None


def profile_field_keys() -> list[str]:
    return [
        "bank_name",
        "branch_name",
        "customer_name",
        "customer_address",
        "account_number",
        "account_type",
        "member_id",
        "currency",
        "reference_no",
        "opening_date",
        "start_date",
        "end_date",
        "opening_balance",
        "target_closing_balance",
        "prepend_statement_mode",
        "prepend_start_date",
        "prepend_anchor_date",
        "prepend_anchor_balance",
        "deposit_min_amount",
        "deposit_max_amount",
        "withdrawal_min_amount",
        "withdrawal_max_amount",
        "amount_rounding_mode",
        "amount_rounding_percentages",
        "interest_rate",
        "tax_rate",
        "cheque_start",
        "include_cheque_column",
        "date_column_mode",
        "closing_row_mode",
        "statement_row_mode",
        "statement_row_count",
        "transaction_count_mode",
        "monthly_transaction_counts",
        "seed",
        "deposit_text",
        "withdrawal_text",
        "description_extra_text",
        "interest_text",
        "tax_text",
        "deposit_mode",
        "withdrawal_mode",
        "first_date_description",
        "last_date_description",
        "deposit_names",
        "withdrawal_names",
    ]


def default_form_values() -> dict[str, str]:
    today = date.today().isoformat()
    return {
        "bank_name": "Nepal Bank Limited",
        "branch_name": "Main Branch, Kathmandu",
        "customer_name": "Customer Name",
        "customer_address": "Customer Address",
        "account_number": "0000000000",
        "account_type": "Saving Account",
        "member_id": "",
        "currency": "NPR",
        "reference_no": "",
        "opening_date": "2022-05-10",
        "start_date": "2025-01-14",
        "end_date": "2026-01-14",
        "opening_balance": "1500000",
        "target_closing_balance": "2550000",
        "prepend_statement_mode": "No",
        "prepend_start_date": "",
        "prepend_anchor_date": "",
        "prepend_anchor_balance": "",
        "deposit_min_amount": "15000",
        "deposit_max_amount": "99000",
        "withdrawal_min_amount": "15000",
        "withdrawal_max_amount": "65000",
        "amount_rounding_mode": "automatic",
        "amount_rounding_percentages": '{"1000":35,"500":35,"100":10,"50":10,"10":0,"5":10}',
        "interest_rate": "8",
        "tax_rate": "6",
        "cheque_start": "10000001",
        "include_cheque_column": "Yes",
        "date_column_mode": "single",
        "first_date_description": "Opening Balance",
        "last_date_description": "Balance C/F",
        "closing_row_mode": "description_only",
        "statement_row_mode": "auto",
        "statement_row_count": "",
        "transaction_count_mode": "auto",
        "monthly_transaction_counts": "",
        "seed": "",
        "deposit_text": "Cash Deposit",
        "withdrawal_text": "Cheque Withdrawal",
        "description_extra_text": "",
        "interest_text": "Interest Posted",
        "tax_text": "Tax Deducted",
        "deposit_mode": "Label + Name",
        "withdrawal_mode": "Label + Name",
        "template_dir": str(DEFAULT_TEMPLATE_DIR),
        "manual_rate": "",
        "rate_mode": "Auto (NRB)",
        "rate_type": "sell",
        "deposit_names": DEFAULT_DEPOSIT_NAMES,
        "withdrawal_names": DEFAULT_WITHDRAWAL_NAMES,
        "today": today,
    }


def auth_store() -> AuthStore:
    global AUTH_STORE
    if AUTH_STORE is None:
        AUTH_STORE = AuthStore(AUTH_DB_FILE)
    return AUTH_STORE


def bootstrap_defaults(current_user: AuthUser) -> dict[str, str]:
    defaults = default_form_values()
    saved_profile = auth_store().load_profile(current_user).get("profile", {})
    if not isinstance(saved_profile, dict):
        return defaults
    for key in profile_field_keys():
        if key in saved_profile:
            defaults[key] = str(saved_profile[key])
    return defaults


def _clean_date_strings(values: object, saturday_only: bool = False) -> set[str]:
    cleaned: set[str] = set()
    if not isinstance(values, list):
        return cleaned
    for item in values:
        text = str(item).strip()
        if not text:
            continue
        try:
            parsed = parse_iso_date(text)
        except Exception:
            continue
        if saturday_only and parsed.weekday() != 5:
            continue
        cleaned.add(parsed.isoformat())
    return cleaned


def _persistent_meta_defaults() -> dict[str, str]:
    return {key: "" for key in PERSISTENT_META_KEYS}


def user_state_file(user_id: int | None = None) -> Path:
    if user_id and user_id > 0:
        USER_STATE_ROOT.mkdir(parents=True, exist_ok=True)
        return USER_STATE_ROOT / f"user_{user_id}.json"
    return STATE_FILE


def _read_state_payload(user_id: int | None = None) -> dict[str, object]:
    state_file = user_state_file(user_id)
    if not state_file.exists() and user_id and STATE_FILE.exists():
        state_file = STATE_FILE
    if not state_file.exists():
        return {}
    try:
        payload = json.loads(state_file.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def _load_existing_persistent_meta(user_id: int | None = None) -> dict[str, str]:
    payload = _read_state_payload(user_id)
    if not payload:
        return _persistent_meta_defaults()
    return {key: str(payload.get(key, "") or "") for key in PERSISTENT_META_KEYS}


def _utc_now_text() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat()


def _parse_meta_datetime(value: object) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _should_auto_refresh(rules: dict[str, object]) -> bool:
    now = datetime.utcnow()
    last_success = _parse_meta_datetime(rules.get("last_auto_refresh_success_at"))
    if last_success is not None and now - last_success < AUTO_REFRESH_SUCCESS_INTERVAL:
        return False
    last_attempt = _parse_meta_datetime(rules.get("last_auto_refresh_attempt_at"))
    if last_attempt is not None and now - last_attempt < AUTO_REFRESH_FAILURE_BACKOFF:
        return False
    return True


def load_persistent_rules(user_id: int | None = None) -> dict[str, object]:
    custom_holidays = set(manual_holiday_dates())
    excluded_saturdays: set[str] = set()
    quarter_date_overrides: dict[str, str] = {}
    synced_quarter_dates: dict[str, str] = {}
    schema_version = 0
    metadata = _persistent_meta_defaults()

    payload = _read_state_payload(user_id)
    if payload:
        schema_version = int(payload.get("schema_version", 1))
        metadata = {key: str(payload.get(key, "") or "") for key in PERSISTENT_META_KEYS}
        if "custom_holidays" in payload:
            custom_holidays = _clean_date_strings(payload.get("custom_holidays", []))
        elif payload.get("holidays"):
            legacy_lines = [item.strip() for item in str(payload.get("holidays", "")).splitlines() if item.strip()]
            custom_holidays = _clean_date_strings(legacy_lines)
        quarter_date_overrides = _clean_quarter_overrides(payload.get("quarter_date_overrides", {}))
        synced_quarter_dates = _clean_quarter_overrides(payload.get("synced_quarter_dates", {}))

    # Seed once so manual deletions survive later loads. Retain saved holidays
    # because older files cannot distinguish manual entries from previous syncs.
    if payload.get("holiday_seed_version") != MANUAL_HOLIDAY_SEED_VERSION:
        custom_holidays.update(manual_holiday_dates())
    if schema_version < PERSISTENT_RULES_SCHEMA_VERSION or payload.get("holiday_seed_version") != MANUAL_HOLIDAY_SEED_VERSION or payload.get("excluded_saturdays"):
        save_persistent_rules(custom_holidays, excluded_saturdays, quarter_date_overrides, synced_quarter_dates, user_id=user_id)
    return {
        "schema_version": PERSISTENT_RULES_SCHEMA_VERSION,
        "custom_holidays": custom_holidays,
        "excluded_saturdays": excluded_saturdays,
        "quarter_date_overrides": quarter_date_overrides,
        "synced_quarter_dates": synced_quarter_dates,
        **metadata,
    }


def save_persistent_rules(
    custom_holidays: set[str],
    excluded_saturdays: set[str],
    quarter_date_overrides: dict[str, str],
    synced_quarter_dates: dict[str, str] | None = None,
    metadata: dict[str, str] | None = None,
    user_id: int | None = None,
) -> None:
    persistent_meta = _load_existing_persistent_meta(user_id)
    if metadata:
        for key in PERSISTENT_META_KEYS:
            if key in metadata:
                persistent_meta[key] = str(metadata.get(key, "") or "")
    payload = {
        "schema_version": PERSISTENT_RULES_SCHEMA_VERSION,
        "holiday_seed_version": MANUAL_HOLIDAY_SEED_VERSION,
        "custom_holidays": sorted(custom_holidays),
        "excluded_saturdays": [],
        "quarter_date_overrides": dict(sorted(_clean_quarter_overrides(quarter_date_overrides).items())),
        "synced_quarter_dates": dict(sorted(_clean_quarter_overrides(synced_quarter_dates or {}).items())),
        **persistent_meta,
    }
    with STATE_LOCK:
        user_state_file(user_id).write_text(json.dumps(payload, indent=2), encoding="utf-8")


def parse_period(start_text: str, end_text: str) -> tuple[date, date] | None:
    try:
        start = parse_iso_date(start_text)
        end = parse_iso_date(end_text)
    except Exception:
        return None
    if start > end:
        return None
    return start, end


def auto_saturday_strings(start_text: str, end_text: str, excluded_saturdays: set[str]) -> list[str]:
    period = parse_period(start_text, end_text)
    if period is None:
        return []
    start, end = period
    current = start
    end = end + timedelta(days=31)
    rows: list[str] = []
    while current <= end:
        iso = current.isoformat()
        if current.weekday() == 5:
            rows.append(iso)
        current += timedelta(days=1)
    return rows


def auto_sunday_holiday_strings(start_text: str, end_text: str) -> list[str]:
    period = parse_period(start_text, end_text)
    if period is None:
        return []
    start, end = period
    if start < SUNDAY_HOLIDAY_START:
        start = SUNDAY_HOLIDAY_START
    current = start
    end = end + timedelta(days=31)
    rows: list[str] = []
    while current <= end:
        if current.weekday() == 6:
            rows.append(current.isoformat())
        current += timedelta(days=1)
    return rows


def blocked_rule_rows(
    start_text: str,
    end_text: str,
    custom_holidays: set[str],
    excluded_saturdays: set[str],
    view: str = "All",
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    sunday_dates = set(auto_sunday_holiday_strings(start_text, end_text))
    saturday_dates = set(auto_saturday_strings(start_text, end_text, excluded_saturdays))
    for holiday in sorted(custom_holidays):
        if is_recurring_holiday(parse_iso_date(holiday)):
            continue
        rows.append({"id": f"holiday:{holiday}", "date": holiday, "type": "Holiday"})
    for holiday in sorted(sunday_dates):
        rows.append({"id": f"sunday:{holiday}", "date": holiday, "type": "Sunday"})
    for saturday in sorted(saturday_dates):
        rows.append({"id": f"saturday:{saturday}", "date": saturday, "type": "Saturday"})

    if view == "Holiday":
        rows = [row for row in rows if row["type"] == "Holiday"]
    elif view == "Sunday":
        rows = [row for row in rows if row["type"] == "Sunday"]
    elif view == "Saturday":
        rows = [row for row in rows if row["type"] == "Saturday"]
    return rows


def blocked_dates(start_text: str, end_text: str, custom_holidays: set[str], excluded_saturdays: set[str]) -> set[date]:
    blocked = {parse_iso_date(item) for item in custom_holidays}
    blocked.update(parse_iso_date(item) for item in auto_sunday_holiday_strings(start_text, end_text))
    blocked.update(parse_iso_date(item) for item in auto_saturday_strings(start_text, end_text, excluded_saturdays))
    return blocked


def build_holiday_payload(start_text: str, end_text: str, view: str = "All", user_id: int | None = None) -> dict[str, object]:
    rules = load_persistent_rules(user_id)
    custom_holidays = rules["custom_holidays"]
    excluded_saturdays = rules["excluded_saturdays"]
    rows = blocked_rule_rows(start_text, end_text, custom_holidays, excluded_saturdays, view=view)
    period = parse_period(start_text, end_text)
    return {
        "view": view,
        "rows": rows,
        "period": {
            "start_date": period[0].isoformat() if period else "",
            "end_date": period[1].isoformat() if period else "",
        },
        "counts": {
            "holidays": len(blocked_rule_rows(start_text, end_text, custom_holidays, excluded_saturdays, view="Holiday")),
            "sundays": len(blocked_rule_rows(start_text, end_text, custom_holidays, excluded_saturdays, view="Sunday")),
            "saturdays": len(blocked_rule_rows(start_text, end_text, custom_holidays, excluded_saturdays, view="Saturday")),
            "showing": len(rows),
        },
    }


def _validate_rule_date(date_text: str, rule_type: str) -> str:
    parsed = parse_iso_date(date_text)
    normalized_type = rule_type.strip().title()
    if normalized_type == "Saturday" and parsed.weekday() != 5:
        raise ValueError("Selected Saturday date must actually be a Saturday.")
    if normalized_type == "Sunday":
        if parsed.weekday() != 6:
            raise ValueError("Selected Sunday date must actually be a Sunday.")
        if parsed < SUNDAY_HOLIDAY_START:
            raise ValueError(f"Automatic Sunday holidays start from {SUNDAY_HOLIDAY_START.isoformat()}.")
    if normalized_type not in {"Holiday", "Saturday", "Sunday"}:
        raise ValueError("Rule type must be Holiday, Sunday, or Saturday.")
    return parsed.isoformat()


def _clean_quarter_overrides(values: object) -> dict[str, str]:
    cleaned: dict[str, str] = {}
    if not isinstance(values, dict):
        return cleaned
    for period_key, date_text in values.items():
        try:
            normalized_key = _validate_quarter_period_key(str(period_key))
            parsed = parse_iso_date(str(date_text))
        except Exception:
            continue
        if parsed.strftime("%Y-%m") != normalized_key:
            continue
        cleaned[normalized_key] = parsed.isoformat()
    return cleaned


def _validate_quarter_period_key(period_key: str) -> str:
    period_key = period_key.strip()
    try:
        parsed = datetime.strptime(period_key, "%Y-%m")
    except ValueError as error:
        raise ValueError("Select a valid quarter period first.") from error
    if parsed.month not in QUARTER_MONTHS:
        raise ValueError("Interest and tax dates can only be set for quarter months.")
    return parsed.strftime("%Y-%m")


def _posting_period_range(start_text: str, end_text: str) -> tuple[int, int]:
    period = parse_period(start_text, end_text)
    if period is None:
        today = date.today()
        return today.year - 1, today.year + 25
    start, end = period
    return start.year - 1, end.year + 25


def _quarter_override_lookup(values: dict[str, str]) -> dict[tuple[int, int], date]:
    lookup: dict[tuple[int, int], date] = {}
    for period_key, date_text in values.items():
        try:
            year_text, month_text = period_key.split("-", 1)
            lookup[(int(year_text), int(month_text))] = parse_iso_date(date_text)
        except Exception:
            continue
    return lookup


def merged_quarter_override_map(start_text: str, end_text: str, user_id: int | None = None) -> dict[str, str]:
    rules = load_persistent_rules(user_id)
    merged = dict(rules["synced_quarter_dates"])
    merged.update(dict(rules["quarter_date_overrides"]))
    return _clean_quarter_overrides(merged)


def build_posting_date_payload(start_text: str, end_text: str, user_id: int | None = None) -> dict[str, object]:
    rules = load_persistent_rules(user_id)
    custom_overrides = dict(rules["quarter_date_overrides"])
    synced_overrides = dict(rules["synced_quarter_dates"])
    override_lookup = _quarter_override_lookup(merged_quarter_override_map(start_text, end_text, user_id))
    start_year, end_year = _posting_period_range(start_text, end_text)
    rows: list[dict[str, str]] = []
    for year in range(start_year, end_year + 1):
        for month in QUARTER_MONTHS:
            posting_date = posting_date_for_period(year, month, override_lookup)
            period_key = f"{year:04d}-{month:02d}"
            rows.append(
                {
                    "period_key": period_key,
                    "label": posting_date.strftime("%b %Y"),
                    "date": posting_date.isoformat(),
                    "source": "Custom" if period_key in custom_overrides else ("Auto" if period_key in synced_overrides else "Default"),
                }
            )
    return {
        "rows": rows,
        "period_options": [{"period_key": row["period_key"], "label": row["label"]} for row in rows],
        "selected_period_key": rows[0]["period_key"] if rows else "",
    }


def _remove_rule(custom_holidays: set[str], excluded_saturdays: set[str], rule_type: str, date_text: str) -> None:
    normalized_date = _validate_rule_date(date_text, rule_type)
    normalized_type = rule_type.strip().title()
    if is_recurring_holiday(parse_iso_date(normalized_date)) or normalized_type in {"Saturday", "Sunday"}:
        raise ValueError(f"Automatic {normalized_type} holidays cannot be removed or modified.")
    custom_holidays.discard(normalized_date)


def _apply_rule(custom_holidays: set[str], excluded_saturdays: set[str], date_text: str, rule_type: str) -> None:
    normalized = _validate_rule_date(date_text, rule_type)
    normalized_type = rule_type.strip().title()
    if normalized_type != "Holiday" or is_recurring_holiday(parse_iso_date(normalized)):
        raise ValueError("Saturday and automatic Sunday holidays do not need a manual rule.")
    custom_holidays.add(normalized)


def apply_rule_action(action_payload: dict[str, object], current_user: AuthUser) -> dict[str, object]:
    rules = load_persistent_rules(current_user.id)
    custom_holidays = set(rules["custom_holidays"])
    excluded_saturdays = set(rules["excluded_saturdays"])
    password = str(action_payload.get("password", "")).strip()
    if not password or not auth_store().verify_user_password(current_user.id, password):
        raise PermissionError("Enter your account password to change manual holidays.")

    action = str(action_payload.get("action", "")).strip()
    if action == "restore_saturdays":
        excluded_saturdays.clear()
    elif action == "add":
        _apply_rule(
            custom_holidays,
            excluded_saturdays,
            str(action_payload.get("date", "")).strip(),
            str(action_payload.get("type", "Holiday")).strip(),
        )
    elif action == "update":
        _remove_rule(
            custom_holidays,
            excluded_saturdays,
            str(action_payload.get("original_type", "Holiday")).strip(),
            str(action_payload.get("original_date", "")).strip(),
        )
        _apply_rule(
            custom_holidays,
            excluded_saturdays,
            str(action_payload.get("date", "")).strip(),
            str(action_payload.get("type", "Holiday")).strip(),
        )
    elif action == "delete":
        _remove_rule(
            custom_holidays,
            excluded_saturdays,
            str(action_payload.get("type", "Holiday")).strip(),
            str(action_payload.get("date", "")).strip(),
        )
    else:
        raise ValueError("Unknown holiday action.")

    save_persistent_rules(custom_holidays, excluded_saturdays, dict(rules["quarter_date_overrides"]), dict(rules["synced_quarter_dates"]), user_id=current_user.id)
    return build_holiday_payload(
        str(action_payload.get("start_date", default_form_values()["start_date"])),
        str(action_payload.get("end_date", default_form_values()["end_date"])),
        view=str(action_payload.get("view", "All")) or "All",
        user_id=current_user.id,
    )


def apply_posting_date_action(action_payload: dict[str, object], current_user: AuthUser) -> dict[str, object]:
    rules = load_persistent_rules(current_user.id)
    custom_holidays = set(rules["custom_holidays"])
    excluded_saturdays = set(rules["excluded_saturdays"])
    quarter_date_overrides = dict(rules["quarter_date_overrides"])
    password = str(action_payload.get("password", "")).strip()
    if not password or not auth_store().verify_user_password(current_user.id, password):
        raise PermissionError("Enter your account password to change interest and tax dates.")

    action = str(action_payload.get("action", "")).strip()
    period_key = _validate_quarter_period_key(str(action_payload.get("period_key", "")).strip())
    if action == "delete":
        quarter_date_overrides.pop(period_key, None)
    elif action in {"add", "update"}:
        posting_date = parse_iso_date(str(action_payload.get("date", "")).strip())
        if posting_date.strftime("%Y-%m") != period_key:
            raise ValueError("Interest and tax posting date must stay inside the selected quarter month.")
        quarter_date_overrides[period_key] = posting_date.isoformat()
    else:
        raise ValueError("Unknown interest and tax date action.")

    save_persistent_rules(custom_holidays, excluded_saturdays, quarter_date_overrides, dict(rules["synced_quarter_dates"]), user_id=current_user.id)
    return build_posting_date_payload(
        str(action_payload.get("start_date", default_form_values()["start_date"])),
        str(action_payload.get("end_date", default_form_values()["end_date"])),
        user_id=current_user.id,
    )


def _download_hamropatro_page(url: str, timeout: int = MANUAL_HAMROPATRO_TIMEOUT) -> str:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) StatementGenerator/2.0",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:  # noqa: S310
        return response.read().decode("utf-8", errors="ignore")


def _quarter_bs_calendar_for_period(year: int, month: int) -> tuple[int, int]:
    mapping = {1: (year + 56, 9), 4: (year + 56, 12), 7: (year + 57, 3), 10: (year + 57, 6)}
    if month not in mapping:
        raise ValueError("Quarter month mapping is not available.")
    return mapping[month]


def _extract_bs_month_last_ad_date(html: str, bs_year: int, bs_month: int) -> str | None:
    pattern = re.compile(rf"editNotes\('([0-9]{{4}}-[0-9]{{1,2}}-[0-9]{{1,2}})','{bs_year}-{bs_month}-([0-9]{{1,2}})'\)")
    best_day = -1
    best_date: str | None = None
    for match in pattern.finditer(html):
        try:
            day = int(match.group(2))
            resolved = datetime.fromisoformat(match.group(1)).date().isoformat()
        except Exception:
            continue
        if day > best_day:
            best_day = day
            best_date = resolved
    if best_date:
        return best_date

    month_label = BS_MONTH_LABELS.get(bs_month, "")
    if not month_label:
        return None
    fallback_pattern = re.compile(
        rf"\b(\d{{1,2}})\s+{re.escape(month_label)}\s+{bs_year}\b.*?([A-Za-z]{{3,9}}\s+\d{{1,2}},\s+\d{{4}})",
        flags=re.IGNORECASE | re.DOTALL,
    )
    for match in fallback_pattern.finditer(html):
        try:
            day = int(match.group(1))
            resolved = datetime.strptime(match.group(2), "%B %d, %Y").date().isoformat()
        except Exception:
            continue
        if day > best_day:
            best_day = day
            best_date = resolved
    return best_date


def _quarter_sync_periods(start_text: str, end_text: str, future_years: int) -> list[tuple[int, int]]:
    period = parse_period(start_text, end_text)
    if period is None:
        today = date.today()
        window_start = today - timedelta(days=120)
        window_end = today + timedelta(days=365 * max(1, future_years))
    else:
        start, end = period
        window_start = start - timedelta(days=120)
        window_end = end + timedelta(days=365 * max(0, future_years))

    periods: list[tuple[int, int]] = []
    for year in range(window_start.year - 1, window_end.year + 2):
        for month in QUARTER_MONTHS:
            candidate = posting_date_for_period(year, month)
            if window_start <= candidate <= window_end:
                periods.append((year, month))
    return sorted(set(periods))


def _compute_hamropatro_quarter_dates(
    start_text: str,
    end_text: str,
    future_years: int = 4,
    timeout: int = MANUAL_HAMROPATRO_TIMEOUT,
) -> dict[str, str]:
    periods = _quarter_sync_periods(start_text, end_text, future_years)
    page_cache: dict[str, str] = {}
    results: dict[str, str] = {}
    for year, month in periods:
        bs_year, bs_month = _quarter_bs_calendar_for_period(year, month)
        cache_key = f"{bs_year}-{bs_month}"
        if cache_key not in page_cache:
            try:
                page_cache[cache_key] = _download_hamropatro_page(
                    f"{HAMROPATRO_ENGLISH_CALENDAR_URL}{bs_year}/{bs_month}",
                    timeout=timeout,
                )
            except Exception:
                page_cache[cache_key] = ""
        if not page_cache[cache_key]:
            continue
        resolved = _extract_bs_month_last_ad_date(page_cache[cache_key], bs_year, bs_month)
        if resolved:
            results[f"{year:04d}-{month:02d}"] = resolved
    return dict(sorted(results.items()))


def sync_hamropatro_posting_dates(action_payload: dict[str, object], current_user: AuthUser) -> dict[str, object]:
    password = str(action_payload.get("password", "")).strip()
    if not password or not auth_store().verify_user_password(current_user.id, password):
        raise PermissionError("Enter your account password to update interest and tax dates automatically.")
    rules = load_persistent_rules(current_user.id)
    synced = _compute_hamropatro_quarter_dates(
        str(action_payload.get("start_date", default_form_values()["start_date"])),
        str(action_payload.get("end_date", default_form_values()["end_date"])),
        future_years=4,
        timeout=MANUAL_HAMROPATRO_TIMEOUT,
    )
    rules = load_persistent_rules(current_user.id)
    save_persistent_rules(
        set(rules["custom_holidays"]),
        set(rules["excluded_saturdays"]),
        dict(rules["quarter_date_overrides"]),
        synced,
        user_id=current_user.id,
    )
    payload = build_posting_date_payload(
        str(action_payload.get("start_date", default_form_values()["start_date"])),
        str(action_payload.get("end_date", default_form_values()["end_date"])),
        user_id=current_user.id,
    )
    payload["sync"] = {
        "detected_count": len(synced),
        "source": HAMROPATRO_ENGLISH_CALENDAR_URL,
        "warning": "" if synced else "Hamro Patro did not return fresh quarter dates right now, so the last saved dates are still being used.",
    }
    return payload


def refresh_from_internet(start_text: str, end_text: str, user_id: int | None = None) -> dict[str, object]:
    rules = load_persistent_rules(user_id)
    if not _should_auto_refresh(rules):
        return {
            "holiday_dates": sorted(set(rules["custom_holidays"])),
            "synced_quarter_dates": dict(rules["synced_quarter_dates"]),
            "skipped": True,
        }
    custom_holidays = set(rules["custom_holidays"])
    excluded_saturdays = set(rules["excluded_saturdays"])
    quarter_date_overrides = dict(rules["quarter_date_overrides"])
    metadata = {
        "last_auto_refresh_attempt_at": _utc_now_text(),
        "last_auto_refresh_success_at": str(rules.get("last_auto_refresh_success_at", "") or ""),
    }
    refreshed = False
    synced_quarter_dates = dict(rules["synced_quarter_dates"])
    try:
        detected_quarter_dates = _compute_hamropatro_quarter_dates(
            start_text,
            end_text,
            future_years=0,
            timeout=AUTO_HAMROPATRO_CALENDAR_TIMEOUT,
        )
        if detected_quarter_dates:
            synced_quarter_dates = detected_quarter_dates
            refreshed = True
    except Exception:
        synced_quarter_dates = dict(rules["synced_quarter_dates"])
    if refreshed:
        metadata["last_auto_refresh_success_at"] = metadata["last_auto_refresh_attempt_at"]
    latest_rules = load_persistent_rules(user_id)
    custom_holidays = set(latest_rules["custom_holidays"])
    quarter_date_overrides = dict(latest_rules["quarter_date_overrides"])
    save_persistent_rules(
        custom_holidays,
        excluded_saturdays,
        quarter_date_overrides,
        synced_quarter_dates,
        metadata=metadata,
        user_id=user_id,
    )
    return {
        "holiday_dates": sorted(custom_holidays),
        "synced_quarter_dates": synced_quarter_dates,
        "skipped": False,
    }


def _admin_restore_target_user_id(action_payload: dict[str, object]) -> int:
    raw_user_id = str(action_payload.get("user_id", "")).strip()
    if not raw_user_id:
        raise ValueError("Select a user before restoring date rules.")
    try:
        target_user_id = int(raw_user_id)
    except ValueError as error:
        raise ValueError("Select a valid user before restoring date rules.") from error
    known_ids = {int(row.get("id", 0) or 0) for row in auth_store().list_users()}
    if target_user_id not in known_ids:
        raise ValueError("Selected user was not found.")
    return target_user_id


def restore_user_date_rules(action_payload: dict[str, object], current_user: AuthUser) -> dict[str, object]:
    if not current_user.is_admin:
        raise PermissionError("Only the admin account can restore user date rules.")
    password = str(action_payload.get("password", "")).strip()
    if not password or not auth_store().verify_user_password(current_user.id, password):
        raise PermissionError("Enter your admin password to restore date rules.")
    target_user_id = _admin_restore_target_user_id(action_payload)
    part = str(action_payload.get("part", "all")).strip().lower()
    if part not in {"all", "holidays", "saturdays", "sundays", "posting_dates"}:
        raise ValueError("Restore type must be holidays, saturdays, sundays, posting_dates, or all.")

    rules = load_persistent_rules(target_user_id)
    custom_holidays = set(rules["custom_holidays"])
    excluded_saturdays = set(rules["excluded_saturdays"])
    quarter_date_overrides = dict(rules["quarter_date_overrides"])
    synced_quarter_dates = dict(rules["synced_quarter_dates"])
    metadata = _persistent_meta_defaults()

    if part in {"all", "holidays"}:
        custom_holidays = set(manual_holiday_dates())
    if part in {"all", "saturdays"}:
        excluded_saturdays = set()
    if part in {"all", "posting_dates"}:
        quarter_date_overrides = {}
        synced_quarter_dates = {}
    save_persistent_rules(
        custom_holidays,
        excluded_saturdays,
        quarter_date_overrides,
        synced_quarter_dates,
        metadata=metadata,
        user_id=target_user_id,
    )

    start_date = str(action_payload.get("start_date", default_form_values()["start_date"]))
    end_date = str(action_payload.get("end_date", default_form_values()["end_date"]))
    return {
        "restored": True,
        "user_id": target_user_id,
        "part": part,
        "holidays": build_holiday_payload(start_date, end_date, view="All", user_id=target_user_id),
        "posting_dates": build_posting_date_payload(start_date, end_date, user_id=target_user_id),
    }


def _custom_template_dir(kind: str, user_id: int | None = None) -> Path:
    normalized = kind.strip().lower()
    if normalized not in {"statement", "certificate"}:
        raise ValueError("Template kind must be statement or certificate.")
    owner_key = f"user_{user_id}" if user_id and user_id > 0 else "shared"
    path = CUSTOM_TEMPLATE_ROOT / owner_key / normalized
    path.mkdir(parents=True, exist_ok=True)
    return path


def _template_owner_dir(user_id: int | None = None) -> Path:
    owner_key = f"user_{user_id}" if user_id and user_id > 0 else "shared"
    path = CUSTOM_TEMPLATE_ROOT / owner_key
    path.mkdir(parents=True, exist_ok=True)
    return path


def _hidden_template_path(user_id: int | None = None) -> Path:
    return _template_owner_dir(user_id) / HIDDEN_TEMPLATE_FILE


def _hidden_template_keys(user_id: int | None = None) -> set[str]:
    path = _hidden_template_path(user_id)
    if not path.exists():
        return set()
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return set()
    if isinstance(raw, list):
        return {str(item).strip().lower() for item in raw if str(item).strip()}
    return set()


def _write_hidden_template_keys(keys: set[str], user_id: int | None = None) -> None:
    path = _hidden_template_path(user_id)
    path.write_text(json.dumps(sorted(keys), indent=2), encoding="utf-8")


def _template_meta_path(path: Path) -> Path:
    return path.with_suffix(path.suffix + TEMPLATE_META_SUFFIX)


def _read_template_meta(path: Path) -> dict[str, object]:
    meta_path = _template_meta_path(path)
    if not meta_path.exists():
        return {}
    try:
        payload = json.loads(meta_path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def _write_template_meta(path: Path, meta: dict[str, object]) -> None:
    _template_meta_path(path).write_text(json.dumps(meta, indent=2, ensure_ascii=True, default=json_default), encoding="utf-8")


def _coerce_template_profile(value: object) -> dict[str, str]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except Exception:
            value = {}
    if not isinstance(value, dict):
        return {}
    profile: dict[str, str] = {}
    allowed_keys = set(profile_field_keys()) | {"template_dir", "manual_rate", "rate_mode", "rate_type"}
    for key, raw in value.items():
        if key in allowed_keys and (isinstance(raw, (str, int, float, bool)) or raw is None):
            profile[key] = "" if raw is None else str(raw)
    return profile


def _openpyxl_color_text(color) -> str:
    if color is None:
        return ""
    try:
        if color.type == "rgb" and color.rgb:
            return str(color.rgb)
        if color.type == "indexed" and color.indexed is not None:
            return f"indexed:{color.indexed}"
        if color.type == "theme" and color.theme is not None:
            return f"theme:{color.theme}"
    except Exception:
        return ""
    return ""


def _excel_cell_style(cell) -> dict[str, object]:
    font = cell.font
    fill = cell.fill
    alignment = cell.alignment
    border = cell.border
    border_style = ""
    try:
        for side in (border.left, border.right, border.top, border.bottom):
            side_style = getattr(side, "style", None)
            if side_style:
                border_style = str(side_style)
                break
    except Exception:
        border_style = ""
    border_css = ""
    border_sides: dict[str, str] = {}
    try:
        side_map = {
            "top": border.top,
            "right": border.right,
            "bottom": border.bottom,
            "left": border.left,
        }
        for side_name, side in side_map.items():
            side_style = getattr(side, "style", None)
            if side_style:
                width = "2px" if side_style in {"medium", "thick", "double"} else "1px"
                color = _openpyxl_color(getattr(side, "color", None)) or "#16324a"
                css = f"{width} solid {color}"
                border_sides[f"border_{side_name}"] = css
        if border_sides and len(set(border_sides.values())) == 1 and len(border_sides) == 4:
            border_css = next(iter(border_sides.values()))
    except Exception:
        border_css = ""
    return {
        "number_format": str(cell.number_format or ""),
        "horizontal": str(alignment.horizontal or ""),
        "vertical": str(alignment.vertical or ""),
        "wrap_text": bool(alignment.wrap_text),
        "indent": float(alignment.indent or 0),
        "shrink_to_fit": bool(alignment.shrink_to_fit),
        "bold": bool(font.bold),
        "italic": bool(font.italic),
        "underline": bool(font.underline),
        "strike": bool(font.strike),
        "font_name": str(font.name or ""),
        "font_size": float(font.sz) if font.sz else "",
        "font_color": _openpyxl_color_text(font.color),
        "fill_color": _openpyxl_color_text(fill.fgColor),
        "border_style": border_style,
        "border": border_css,
        **border_sides,
    }


def _excel_cell_runs(cell) -> list[dict[str, object]]:
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
    except Exception:
        return []
    value = cell.value
    if not isinstance(value, CellRichText):
        return []
    runs: list[dict[str, object]] = []
    for part in value:
        if isinstance(part, str):
            if part:
                runs.append({"text": part})
            continue
        if not isinstance(part, TextBlock):
            continue
        text = str(part.text or "")
        if not text:
            continue
        font = part.font
        run: dict[str, object] = {"text": text}
        if getattr(font, "vertAlign", None) == "subscript":
            run["vert_align"] = "subscript"
        elif getattr(font, "vertAlign", None) == "superscript":
            run["vert_align"] = "superscript"
        if getattr(font, "b", None):
            run["bold"] = True
        if getattr(font, "i", None):
            run["italic"] = True
        if getattr(font, "u", None):
            run["underline"] = True
        if getattr(font, "strike", None):
            run["strike"] = True
        if getattr(font, "rFont", None):
            run["font_name"] = str(font.rFont)
        if getattr(font, "sz", None):
            run["font_size"] = float(font.sz)
        if getattr(font, "color", None) is not None:
            color_text = _openpyxl_color_text(font.color)
            if color_text:
                run["font_color"] = color_text
        runs.append(run)
    return runs


def _header_footer_text(section) -> str:
    parts: list[str] = []
    for label, area in (("Left", "left"), ("Center", "center"), ("Right", "right")):
        try:
            text = str(getattr(section, area).text or "")
        except Exception:
            text = ""
        if text:
            parts.append(f"{label}: {text}")
    return "\n".join(parts)


def _normal_alignment_text(value: object) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    if "center" in text or text == "middle":
        return "center"
    if "right" in text or text == "end":
        return "right"
    if "justify" in text or "both" in text:
        return "justify"
    if "left" in text or text == "start":
        return "left"
    return ""


def _excel_horizontal_alignment(value: object) -> str | None:
    text = str(value or "").strip().lower()
    if not text:
        return None
    mapping = {
        "middle": "center",
        "centre": "center",
        "center": "center",
        "left": "left",
        "right": "right",
        "justify": "justify",
        "distributed": "distributed",
        "fill": "fill",
        "general": "general",
        "centercontinuous": "centerContinuous",
        "center continuous": "centerContinuous",
    }
    return mapping.get(text)


def _excel_vertical_alignment(value: object) -> str | None:
    text = str(value or "").strip().lower()
    if not text:
        return None
    mapping = {
        "middle": "center",
        "centre": "center",
        "center": "center",
        "top": "top",
        "bottom": "bottom",
        "justify": "justify",
        "distributed": "distributed",
    }
    return mapping.get(text)


def _apply_excel_page_setup(worksheet) -> None:
    try:
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.35
        worksheet.page_margins.bottom = 0.35
        worksheet.page_margins.header = 0.2
        worksheet.page_margins.footer = 0.2
        max_row, max_column = _worksheet_meaningful_bounds(worksheet)
        if max_row and max_column:
            from openpyxl.utils import get_column_letter  # noqa: E402

            worksheet.print_area = f"A1:{get_column_letter(max_column)}{max_row}"
    except Exception:
        pass


def _cell_has_meaningful_value(cell) -> bool:
    value = getattr(cell, "value", None)
    return value is not None and str(value).strip() != ""


def _worksheet_meaningful_bounds(worksheet) -> tuple[int, int]:
    max_row = 0
    max_column = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if not _cell_has_meaningful_value(cell):
                continue
            max_row = max(max_row, int(cell.row))
            max_column = max(max_column, int(cell.column))
    value_max_column = max_column
    try:
        from openpyxl.utils import range_boundaries  # noqa: E402

        for merged_range in worksheet.merged_cells.ranges:
            min_col, min_row, max_col, merged_max_row = range_boundaries(str(merged_range))
            start_cell = worksheet.cell(row=min_row, column=min_col)
            start_has_value = _cell_has_meaningful_value(start_cell)
            if not (start_has_value or getattr(start_cell, "has_style", False)):
                continue
            max_row = max(max_row, int(merged_max_row))
            if start_has_value:
                max_column = max(max_column, int(max_col))
            elif value_max_column:
                if min_col <= value_max_column:
                    max_column = max(max_column, min(int(max_col), int(value_max_column)))
            else:
                max_column = max(max_column, int(max_col))
    except Exception:
        pass
    if max_row == 0 or max_column == 0:
        return max(1, int(getattr(worksheet, "max_row", 1) or 1)), max(1, int(getattr(worksheet, "max_column", 1) or 1))
    return max_row, max_column


def _apply_word_page_setup(document) -> None:
    try:
        from docx.shared import Mm  # noqa: E402

        for section in document.sections:
            section.page_width = Mm(210)
            section.page_height = Mm(297)
            section.left_margin = Mm(12.7)
            section.right_margin = Mm(12.7)
            section.top_margin = Mm(12.7)
            section.bottom_margin = Mm(12.7)
    except Exception:
        pass


def _certificate_object_key(text: str) -> str:
    lower = " ".join(str(text or "").lower().split())
    if not lower:
        return ""
    if "ref" in lower or "reference" in lower:
        return "reference_no_date" if "date" in lower else "reference_no"
    if re.search(r"\b(issue\s*)?date\b", lower):
        return "issue_date_slash"
    if re.search(r"\b(name|account holder|customer)\b", lower):
        return "customer_name"
    if "address" in lower:
        return "customer_address"
    if re.search(r"\b(a/c|account)\s*(no|number)\b", lower):
        return "account_number_member_id" if "member" in lower else "account_number"
    if "member" in lower:
        return "member_id"
    if "currency" in lower:
        return "currency"
    if "exchange" in lower and "rate" in lower:
        return "usd_npr_text"
    if "usd" in lower and "word" in lower:
        return "balance_words_usd"
    if "usd" in lower:
        return "equivalent_usd_text"
    if "word" in lower:
        return "balance_words_npr"
    if "balance" in lower:
        return "total_balance_npr_text"
    if "authoriz" in lower:
        return "authorization_details"
    return ""


def _certificate_object_label(object_key: str) -> str:
    labels = {
        "reference_no_date": "Ref No. / Date",
        "reference_no": "Ref No.",
        "issue_date_slash": "Date",
        "customer_name": "Name",
        "customer_address": "Address",
        "account_number_member_id": "Account Number / Member ID",
        "account_number": "Account Number",
        "member_id": "Member ID",
        "currency": "Currency",
        "total_balance_npr_text": "Total Balance",
        "balance_words_npr": "Balance In Word Format",
        "usd_npr_text": "Exchange Rate",
        "equivalent_usd_text": "Equivalent in USD Balance",
        "balance_words_usd": "USD Balance in Word",
        "authorization_details": "Authorization Details",
    }
    return labels.get(object_key, object_key)


def _word_run_vert_align(run) -> str:
    try:
        if run.font.subscript:
            return "subscript"
        if run.font.superscript:
            return "superscript"
    except Exception:
        pass
    try:
        vert_align = run._element.rPr.vertAlign
        value = str(vert_align.val or "").lower() if vert_align is not None else ""
        if value == "subscript":
            return "subscript"
        if value == "superscript":
            return "superscript"
    except Exception:
        return ""
    return ""


def _word_run_color_text(run) -> str:
    try:
        rgb = run.font.color.rgb
        return str(rgb) if rgb is not None else ""
    except Exception:
        return ""


def _word_paragraph_runs(paragraphs) -> list[dict[str, object]]:
    runs: list[dict[str, object]] = []
    for paragraph_index, paragraph in enumerate(paragraphs):
        if paragraph_index > 0:
            runs.append({"text": "\n"})
        paragraph_runs = list(paragraph.runs)
        if not paragraph_runs and paragraph.text:
            runs.append({"text": paragraph.text})
            continue
        for run in paragraph_runs:
            text = run.text
            if text == "":
                continue
            font = run.font
            size = getattr(font, "size", None)
            run_payload: dict[str, object] = {"text": text}
            vert_align = _word_run_vert_align(run)
            if vert_align:
                run_payload["vert_align"] = vert_align
            if getattr(font, "bold", None) is True:
                run_payload["bold"] = True
            if getattr(font, "italic", None) is True:
                run_payload["italic"] = True
            if getattr(font, "underline", None):
                run_payload["underline"] = True
            if getattr(font, "strike", None) is True:
                run_payload["strike"] = True
            if getattr(font, "name", None):
                run_payload["font_name"] = str(font.name)
            if size is not None:
                run_payload["font_size"] = round(size.pt, 2)
            color = _word_run_color_text(run)
            if color:
                run_payload["font_color"] = color
            runs.append(run_payload)
    return runs


def _scan_excel_template(path: Path, profile: dict[str, str]) -> dict[str, object]:
    if path.suffix.lower() != ".xlsx":
        return {
            "kind": "statement",
            "file_type": path.suffix.lower(),
            "summary": {
                "file": path.name,
                "active_cell_count": 0,
                "warning": "Legacy .xls files can be exported, but web scan/edit needs .xlsx.",
            },
            "items": [],
            "profile": profile,
        }
    from openpyxl import load_workbook  # noqa: E402
    from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries  # noqa: E402

    workbook = load_workbook(path, data_only=False, rich_text=True)
    items: list[dict[str, object]] = []
    sheets: list[dict[str, object]] = []
    for worksheet in workbook.worksheets:
        render_max_row, render_max_column = _worksheet_meaningful_bounds(worksheet)
        merged = [str(item) for item in worksheet.merged_cells.ranges]
        merged_lookup: dict[str, dict[str, object]] = {}
        covered_merged_cells: set[str] = set()
        for merged_range in worksheet.merged_cells.ranges:
            try:
                min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
            except Exception:
                continue
            if min_row > render_max_row or min_col > render_max_column:
                continue
            max_row = min(max_row, render_max_row)
            max_col = min(max_col, render_max_column)
            start_address = f"{get_column_letter(min_col)}{min_row}"
            merged_lookup[start_address] = {
                "merge_range": str(merged_range),
                "colspan": max_col - min_col + 1,
                "rowspan": max_row - min_row + 1,
            }
            for row_index in range(min_row, max_row + 1):
                for col_index in range(min_col, max_col + 1):
                    address = f"{get_column_letter(col_index)}{row_index}"
                    if address != start_address:
                        covered_merged_cells.add(address)
        column_widths = {
            key: float(dimension.width)
            for key, dimension in worksheet.column_dimensions.items()
            if getattr(dimension, "width", None) and column_index_from_string(key) <= render_max_column
        }
        row_heights = {
            str(key): float(dimension.height)
            for key, dimension in worksheet.row_dimensions.items()
            if getattr(dimension, "height", None) and int(key) <= render_max_row
        }
        sheets.append(
            {
                "name": worksheet.title,
                "max_row": render_max_row,
                "max_column": render_max_column,
                "merged_ranges": merged,
                "column_widths": column_widths,
                "row_heights": row_heights,
            }
        )
        for row in worksheet.iter_rows(min_row=1, max_row=render_max_row, min_col=1, max_col=render_max_column):
            for cell in row:
                if cell.coordinate in covered_merged_cells:
                    continue
                value = cell.value
                has_content = value is not None and str(value) != ""
                merge_info = merged_lookup.get(cell.coordinate, {})
                if not has_content and not cell.has_style and not merge_info:
                    continue
                text = "" if value is None else str(value)
                is_formula = text.startswith("=")
                items.append(
                    {
                        "key": f"{worksheet.title}!{cell.coordinate}",
                        "sheet": worksheet.title,
                        "address": cell.coordinate,
                        "label": f"{worksheet.title}!{cell.coordinate}",
                        "text": text,
                        "formula": text if is_formula else "",
                        "is_formula": is_formula,
                        "runs": _excel_cell_runs(cell),
                        "style": _excel_cell_style(cell),
                        **merge_info,
                    }
                )
        for header_footer_key, label, section in (
            ("__header__", "Excel Header", worksheet.oddHeader),
            ("__footer__", "Excel Footer", worksheet.oddFooter),
        ):
            header_footer_text = _header_footer_text(section)
            if header_footer_text:
                items.append(
                    {
                        "key": f"{worksheet.title}!{header_footer_key}",
                        "sheet": worksheet.title,
                        "address": header_footer_key,
                        "label": f"{worksheet.title} {label}",
                        "text": header_footer_text,
                        "style": {"font_name": "Calibri", "font_size": 11, "wrap_text": True},
                    }
                )
    workbook.close()
    return {
        "kind": "statement",
        "file_type": path.suffix.lower(),
        "summary": {
            "file": path.name,
            "active_cell_count": len(items),
            "sheet_count": len(sheets),
            "sheets": sheets,
        },
        "items": items,
        "profile": profile,
    }


def _word_paragraph_style(paragraph) -> dict[str, object]:
    first_run = paragraph.runs[0] if paragraph.runs else None
    font = first_run.font if first_run is not None else None
    size = getattr(font, "size", None)
    paragraph_format = paragraph.paragraph_format

    def length_text(value) -> str:
        if value is None:
            return ""
        try:
            return f"{round(float(value.pt) / 0.75, 2)}px"
        except Exception:
            return ""

    line_spacing = paragraph_format.line_spacing
    try:
        line_spacing_text: object = round(float(line_spacing), 2) if isinstance(line_spacing, (int, float)) else ""
    except Exception:
        line_spacing_text = ""
    left_tab = ""
    try:
        for tab_stop in paragraph_format.tab_stops:
            left_tab = length_text(tab_stop.position)
            if left_tab:
                break
    except Exception:
        left_tab = ""
    return {
        "alignment": _normal_alignment_text(paragraph.alignment),
        "horizontal": _normal_alignment_text(paragraph.alignment),
        "style": str(paragraph.style.name if paragraph.style is not None else ""),
        "bold": bool(getattr(font, "bold", False)) if font is not None else False,
        "italic": bool(getattr(font, "italic", False)) if font is not None else False,
        "underline": bool(getattr(font, "underline", False)) if font is not None else False,
        "strike": bool(getattr(font, "strike", False)) if font is not None else False,
        "font_name": str(getattr(font, "name", "") or "") if font is not None else "",
        "font_size": round(size.pt, 2) if size is not None else "",
        "left_indent": length_text(paragraph_format.left_indent),
        "right_indent": length_text(paragraph_format.right_indent),
        "first_line_indent": length_text(paragraph_format.first_line_indent),
        "left_tab": left_tab,
        "line_spacing": line_spacing_text,
        "space_before": length_text(paragraph_format.space_before),
        "space_after": length_text(paragraph_format.space_after),
    }


def _scan_word_template(path: Path, profile: dict[str, str]) -> dict[str, object]:
    if path.suffix.lower() != ".docx":
        return {
            "kind": "certificate",
            "file_type": path.suffix.lower(),
            "summary": {
                "file": path.name,
                "text_block_count": 0,
                "warning": "Legacy .doc files can be exported, but web scan/edit needs .docx.",
            },
            "items": [],
            "profile": profile,
        }
    from docx import Document  # noqa: E402

    document = Document(str(path))
    items: list[dict[str, object]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
        if text or paragraph.style:
            object_key = _certificate_object_key(text)
            items.append(
                {
                    "key": f"p:{index}",
                    "label": _certificate_object_label(object_key) if object_key else f"Paragraph {index + 1}",
                    "object_key": object_key,
                    "text": text,
                    "runs": _word_paragraph_runs([paragraph]),
                    "style": _word_paragraph_style(paragraph),
                }
            )
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            previous_key = ""
            for cell_index, cell in enumerate(row.cells):
                text = "\n".join(paragraph.text for paragraph in cell.paragraphs)
                object_key = _certificate_object_key(text)
                inferred_key = object_key or (previous_key if not text.strip() else "")
                if object_key:
                    previous_key = object_key
                style = _word_paragraph_style(cell.paragraphs[0]) if cell.paragraphs else {}
                label = f"Table {table_index + 1} R{row_index + 1} C{cell_index + 1}"
                if inferred_key:
                    label = _certificate_object_label(inferred_key)
                    if not text.strip():
                        label = f"{label} Value"
                items.append(
                    {
                        "key": f"t:{table_index}:r:{row_index}:c:{cell_index}",
                        "label": label,
                        "object_key": inferred_key,
                        "text": text,
                        "runs": _word_paragraph_runs(cell.paragraphs),
                        "style": style,
                    }
                )
    return {
        "kind": "certificate",
        "file_type": path.suffix.lower(),
        "summary": {
            "file": path.name,
            "text_block_count": len(items),
            "table_count": len(document.tables),
        },
        "items": items,
        "profile": profile,
    }


def scan_template_file(kind: str, path: Path, profile: dict[str, str] | None = None) -> dict[str, object]:
    path = editable_office_copy(path)
    normalized_kind = kind.strip().lower()
    template_profile = _coerce_template_profile(profile or _read_template_meta(path).get("profile", {}))
    if normalized_kind == "statement":
        return _scan_excel_template(path, template_profile)
    if normalized_kind == "certificate":
        return _scan_word_template(path, template_profile)
    raise ValueError("Template kind must be statement or certificate.")


def _rescan_and_save_template_meta(kind: str, path: Path, profile: dict[str, str] | None = None) -> dict[str, object]:
    existing = _read_template_meta(path)
    preserved_profile = _coerce_template_profile(profile if profile is not None else existing.get("profile", {}))
    scan = scan_template_file(kind, path, preserved_profile)
    meta = {
        "kind": kind,
        "name": path.stem,
        "suffix": path.suffix.lower(),
        "source": "custom",
        "editable": path.suffix.lower() in {".xlsx", ".docx"},
        "profile": preserved_profile,
        "scan": scan,
        "updated_at": _utc_now_text(),
    }
    _write_template_meta(path, meta)
    return meta


def _template_key(kind: str, name: str) -> str:
    return f"{kind.strip().lower()}:{name.strip().lower()}"


def _annotated_entries(entries: list[TemplateEntry], source: str, editable: bool) -> list[TemplateEntry]:
    return [
        TemplateEntry(
            entry.name,
            entry.path,
            source=source,
            editable=bool(editable and entry.path.suffix.lower() in {".xlsx", ".xls", ".docx", ".doc"}),
        )
        for entry in entries
    ]


def _merge_template_entries(bundled: list[TemplateEntry], custom: list[TemplateEntry]) -> list[TemplateEntry]:
    merged: dict[str, TemplateEntry] = {entry.name.lower(): entry for entry in bundled}
    for entry in custom:
        merged[entry.name.lower()] = entry
    return sorted(merged.values(), key=lambda item: item.name.lower())


def combined_template_catalog(template_dir_text: str, user_id: int | None = None) -> TemplateCatalog:
    directory = Path(template_dir_text.strip() or str(DEFAULT_TEMPLATE_DIR))
    hidden_keys = _hidden_template_keys(user_id)
    bundled_catalog = scan_template_directory(directory)
    bundled_catalog = TemplateCatalog(
        statement_templates=[
            entry for entry in bundled_catalog.statement_templates
            if _template_key("statement", entry.name) not in hidden_keys
        ],
        certificate_templates=[
            entry for entry in bundled_catalog.certificate_templates
            if _template_key("certificate", entry.name) not in hidden_keys
        ],
    )
    custom_statement_catalog = scan_template_directory(_custom_template_dir("statement", user_id))
    custom_certificate_catalog = scan_template_directory(_custom_template_dir("certificate", user_id))
    return TemplateCatalog(
        statement_templates=_merge_template_entries(
            _annotated_entries(bundled_catalog.statement_templates, "bundled", False),
            _annotated_entries(custom_statement_catalog.statement_templates, "custom", True),
        ),
        certificate_templates=_merge_template_entries(
            _annotated_entries(bundled_catalog.certificate_templates, "bundled", False),
            _annotated_entries(custom_certificate_catalog.certificate_templates, "custom", True),
        ),
    )


def _serialize_template_entries(entries: list[TemplateEntry]) -> list[dict[str, object]]:
    serialized: list[dict[str, object]] = []
    for entry in entries:
        meta = _read_template_meta(entry.path) if entry.source == "custom" else {}
        scan = meta.get("scan", {}) if isinstance(meta, dict) else {}
        summary = scan.get("summary", {}) if isinstance(scan, dict) else {}
        serialized.append(
            {
            "name": entry.name,
            "suffix": entry.path.suffix,
            "source": entry.source,
            "editable": entry.editable,
                "source_extension": entry.path.suffix.lower(),
                "scan_summary": summary if isinstance(summary, dict) else {},
                "has_profile": bool(meta.get("profile")) if isinstance(meta, dict) else False,
            }
        )
    return serialized


def serialize_catalog(template_dir_text: str, user_id: int | None = None) -> dict[str, object]:
    directory = Path(template_dir_text.strip() or str(DEFAULT_TEMPLATE_DIR))
    catalog = combined_template_catalog(str(directory), user_id)
    return {
        "template_dir": str(directory),
        "statement_templates": _serialize_template_entries(catalog.statement_templates),
        "certificate_templates": _serialize_template_entries(catalog.certificate_templates),
    }


def save_uploaded_template(
    kind: str,
    filename: str,
    content: bytes,
    custom_name: str,
    user_id: int | None = None,
    profile_payload: dict[str, object] | str | None = None,
) -> dict[str, object]:
    normalized_kind = kind.strip().lower()
    suffix = Path(filename).suffix.lower()
    allowed_extensions = STATEMENT_EXTENSIONS if normalized_kind == "statement" else CERTIFICATE_EXTENSIONS
    if suffix not in allowed_extensions:
        allowed = ", ".join(sorted(allowed_extensions))
        raise ValueError(f"Use one of these template file types for {normalized_kind}: {allowed}.")
    target_dir = _custom_template_dir(normalized_kind, user_id)
    base_name = safe_filename(custom_name.strip() or Path(filename).stem)
    output_path = target_dir / f"{base_name}{suffix}"
    output_path.write_bytes(content)
    meta = _rescan_and_save_template_meta(normalized_kind, output_path, _coerce_template_profile(profile_payload or {}))
    return {
        "kind": normalized_kind,
        "name": base_name,
        "suffix": suffix,
        "source": "custom",
        "editable": suffix in {".xlsx", ".xls", ".docx", ".doc"},
        "scan_summary": dict(meta.get("scan", {})).get("summary", {}) if isinstance(meta.get("scan"), dict) else {},
    }


def delete_uploaded_template(kind: str, name: str, user_id: int | None = None, template_dir_text: str = "") -> str:
    normalized_kind = kind.strip().lower()
    target_dir = _custom_template_dir(normalized_kind, user_id)
    normalized_name = name.strip().lower()
    if not normalized_name:
        raise ValueError("Select a template first.")
    candidates = [path for path in target_dir.iterdir() if path.is_file() and path.stem.lower() == normalized_name]
    if candidates:
        for path in candidates:
            path.unlink(missing_ok=True)
            _template_meta_path(path).unlink(missing_ok=True)
        return "custom"
    directory = Path(template_dir_text.strip() or str(DEFAULT_TEMPLATE_DIR))
    bundled_catalog = scan_template_directory(directory)
    entries = bundled_catalog.statement_templates if normalized_kind == "statement" else bundled_catalog.certificate_templates
    if not any(entry.name.lower() == normalized_name for entry in entries):
        raise ValueError("Selected template was not found.")
    hidden_keys = _hidden_template_keys(user_id)
    hidden_keys.add(_template_key(normalized_kind, name))
    _write_hidden_template_keys(hidden_keys, user_id)
    return "bundled"


def resolve_template_entry(kind: str, template_name: str, template_dir_text: str, user_id: int | None = None) -> TemplateEntry | None:
    catalog = combined_template_catalog(template_dir_text, user_id)
    entries = catalog.statement_templates if kind == "statement" else catalog.certificate_templates
    for entry in entries:
        if entry.name == template_name:
            return entry
    return None


def template_detail(kind: str, template_name: str, template_dir_text: str, user_id: int | None = None) -> dict[str, object]:
    normalized_kind = kind.strip().lower()
    selected = resolve_template_entry(normalized_kind, template_name, template_dir_text, user_id)
    if selected is None:
        raise ValueError("Selected template was not found.")
    selected = _require_custom_template_entry(normalized_kind, template_name, template_dir_text, user_id)
    meta = _read_template_meta(selected.path)
    profile = _coerce_template_profile(meta.get("profile", {}))
    scan = scan_template_file(normalized_kind, selected.path, profile)
    if selected.source == "custom":
        _rescan_and_save_template_meta(normalized_kind, selected.path, profile)
    return {
        "kind": normalized_kind,
        "name": selected.name,
        "suffix": selected.path.suffix.lower(),
        "source": selected.source,
        "editable": bool(selected.editable and selected.source == "custom"),
        "profile": profile,
        "scan": scan,
        "html_preview": (re.search(r'<body>(.*)</body>', word_html(editable_office_copy(selected.path), selected.name, True), re.S).group(1)
                         if normalized_kind == 'certificate' else ''),
    }


def _require_custom_template_entry(kind: str, template_name: str, template_dir_text: str, user_id: int | None = None) -> TemplateEntry:
    selected = resolve_template_entry(kind, template_name, template_dir_text, user_id)
    if selected is None:
        raise ValueError("Selected template was not found.")
    if selected.source != "custom":
        source = selected.path
        directory = _custom_template_dir(kind, user_id)
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / (selected.name + source.suffix)
        shutil.copy2(source, target)
        _rescan_and_save_template_meta(kind, target)
        selected = resolve_template_entry(kind, template_name, template_dir_text, user_id)
    return selected


def save_template_profile(kind: str, template_name: str, template_dir_text: str, user_id: int | None, profile_payload: object) -> dict[str, object]:
    normalized_kind = kind.strip().lower()
    selected = _require_custom_template_entry(normalized_kind, template_name, template_dir_text, user_id)
    meta = _rescan_and_save_template_meta(normalized_kind, selected.path, _coerce_template_profile(profile_payload))
    return {
        "saved": True,
        "kind": normalized_kind,
        "name": selected.name,
        "profile": meta.get("profile", {}),
        "scan_summary": dict(meta.get("scan", {})).get("summary", {}) if isinstance(meta.get("scan"), dict) else {},
    }


def _split_excel_key(key: str) -> tuple[str, str]:
    if "!" not in key:
        raise ValueError("Select a valid Excel cell from the scanned format.")
    sheet_name, address = key.split("!", 1)
    if not sheet_name.strip() or not address.strip():
        raise ValueError("Select a valid Excel cell from the scanned format.")
    return sheet_name, address


def _apply_excel_style(cell, style_payload: dict[str, object]) -> None:
    from copy import copy as copy_style  # noqa: E402
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

    if "number_format" in style_payload:
        cell.number_format = str(style_payload.get("number_format") or "General")

    alignment = copy_style(cell.alignment)
    horizontal = _excel_horizontal_alignment(style_payload.get("horizontal", alignment.horizontal or ""))
    vertical = _excel_vertical_alignment(style_payload.get("vertical", alignment.vertical or ""))
    wrap_text = bool(style_payload.get("wrap_text", alignment.wrap_text))
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap_text,
        text_rotation=alignment.text_rotation,
        shrink_to_fit=alignment.shrink_to_fit,
        indent=float(style_payload.get("indent", alignment.indent or 0) or 0),
    )

    font = copy_style(cell.font)
    font_size_raw = style_payload.get("font_size", font.sz)
    try:
        font_size = float(font_size_raw) if str(font_size_raw or "").strip() else font.sz
    except Exception:
        font_size = font.sz
    font_color = str(style_payload.get("font_color", "") or "").strip().lstrip("#")
    if font_color and re.fullmatch(r"[0-9A-Fa-f]{6,8}", font_color):
        if len(font_color) == 6:
            font_color = f"FF{font_color}"
    else:
        font_color = font.color.rgb if getattr(font.color, "type", "") == "rgb" else None
    font_kwargs = {
        "name": str(style_payload.get("font_name", font.name or "") or font.name or "Calibri"),
        "size": font_size,
        "bold": bool(style_payload.get("bold", font.bold)),
        "italic": bool(style_payload.get("italic", font.italic)),
        "underline": "single" if bool(style_payload.get("underline", bool(font.underline))) else None,
        "strike": bool(style_payload.get("strike", font.strike)),
        "color": font_color,
    }
    vert_align = str(style_payload.get("vert_align", font.vertAlign or "") or "").strip()
    if vert_align in {"subscript", "superscript", "baseline"}:
        font_kwargs["vertAlign"] = vert_align
    cell.font = Font(**font_kwargs)

    fill_color = str(style_payload.get("fill_color", "") or "").strip().lstrip("#")
    if fill_color and re.fullmatch(r"[0-9A-Fa-f]{6,8}", fill_color):
        if len(fill_color) == 6:
            fill_color = f"FF{fill_color}"
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color.upper())

    border_style = str(style_payload.get("border_style", "") or "").strip().lower()
    if border_style:
        if border_style == "none":
            cell.border = Border()
        else:
            medium = Side(style="medium", color="FF16324A")
            thin = Side(style="thin", color="FF16324A")
            if border_style == "medium":
                cell.border = Border(left=medium, right=medium, top=medium, bottom=medium)
            elif border_style in {"thin", "outside"}:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            elif border_style == "top":
                cell.border = Border(left=cell.border.left, right=cell.border.right, top=thin, bottom=cell.border.bottom)
            elif border_style == "bottom":
                cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=thin)
            elif border_style == "left":
                cell.border = Border(left=thin, right=cell.border.right, top=cell.border.top, bottom=cell.border.bottom)
            elif border_style == "right":
                cell.border = Border(left=cell.border.left, right=thin, top=cell.border.top, bottom=cell.border.bottom)


def update_excel_template_item(path: Path, key: str, text: str, style_payload: dict[str, object]) -> None:
    from openpyxl import load_workbook  # noqa: E402
    from openpyxl.styles import Alignment  # noqa: E402
    from copy import copy as copy_style  # noqa: E402
    from openpyxl.utils import range_boundaries  # noqa: E402

    workbook = load_workbook(path, rich_text=True)
    target_key = str(style_payload.get("target_key", "") or text or "")
    if key in {"__insert_row__", "__insert_column__", "__add_header__", "__add_footer__"}:
        sheet_name = workbook.sheetnames[0]
        target_address = "A1"
        if "!" in target_key:
            sheet_name, target_address = _split_excel_key(target_key)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError("Selected sheet was not found in this format.")
        worksheet = workbook[sheet_name]
        position = re.fullmatch(r"[A-Z]+(\d+)", target_address)
        row_number = int(position.group(1)) if position else 1
        col_match = re.fullmatch(r"([A-Z]+)\d+", target_address)
        col_number = 1
        if col_match:
            col_number = 0
            for letter in col_match.group(1):
                col_number = col_number * 26 + (ord(letter) - 64)
        if key == "__insert_row__":
            source_row = max(1, min(row_number, worksheet.max_row))
            worksheet.insert_rows(row_number)
            for column in range(1, worksheet.max_column + 1):
                source_cell = worksheet.cell(row=source_row + 1 if source_row >= row_number else source_row, column=column)
                target_cell = worksheet.cell(row=row_number, column=column)
                if source_cell.has_style:
                    target_cell._style = copy_style(source_cell._style)
                target_cell.number_format = source_cell.number_format
                target_cell.alignment = copy_style(source_cell.alignment)
            worksheet.row_dimensions[row_number].height = worksheet.row_dimensions[source_row].height
        elif key == "__insert_column__":
            source_col = max(1, min(col_number, worksheet.max_column))
            worksheet.insert_cols(col_number)
            for row in range(1, worksheet.max_row + 1):
                source_cell = worksheet.cell(row=row, column=source_col + 1 if source_col >= col_number else source_col)
                target_cell = worksheet.cell(row=row, column=col_number)
                if source_cell.has_style:
                    target_cell._style = copy_style(source_cell._style)
                target_cell.number_format = source_cell.number_format
                target_cell.alignment = copy_style(source_cell.alignment)
            from openpyxl.utils import get_column_letter  # noqa: E402

            worksheet.column_dimensions[get_column_letter(col_number)].width = worksheet.column_dimensions[get_column_letter(source_col)].width
        elif key == "__add_header__":
            worksheet.oddHeader.center.text = text
        elif key == "__add_footer__":
            worksheet.oddFooter.center.text = text
        _apply_excel_page_setup(worksheet)
        workbook.save(path)
        workbook.close()
        return

    sheet_name, address = _split_excel_key(key)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError("Selected sheet was not found in this format.")
    worksheet = workbook[sheet_name]
    if address == "__header__":
        worksheet.oddHeader.center.text = text
        _apply_excel_page_setup(worksheet)
        workbook.save(path)
        workbook.close()
        return
    if address == "__footer__":
        worksheet.oddFooter.center.text = text
        _apply_excel_page_setup(worksheet)
        workbook.save(path)
        workbook.close()
        return
    unmerge_range = str(style_payload.get("unmerge_range", "") or "").strip()
    if unmerge_range and re.fullmatch(r"[A-Z]+\d+:[A-Z]+\d+", unmerge_range):
        try:
            worksheet.unmerge_cells(unmerge_range)
        except ValueError:
            pass
    cell = worksheet[address]

    def excel_rich_text_value(runs_payload: object):
        if not isinstance(runs_payload, list) or not runs_payload:
            return None
        from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
        from openpyxl.cell.text import InlineFont  # noqa: E402

        rich_text = CellRichText()
        wrote = False
        for segment in runs_payload:
            if not isinstance(segment, dict):
                continue
            segment_text = str(segment.get("text", ""))
            if segment_text == "":
                continue
            wrote = True
            vert_align = str(segment.get("vert_align", "") or "").strip().lower()
            has_inline_style = bool(
                vert_align
                or segment.get("bold") is True
                or segment.get("italic") is True
                or segment.get("underline") is True
                or segment.get("strike") is True
                or str(segment.get("font_name", "") or "").strip()
                or str(segment.get("font_size", "") or "").strip()
                or str(segment.get("font_color", "") or "").strip()
            )
            if not has_inline_style:
                rich_text.append(segment_text)
                continue
            font_color = str(segment.get("font_color", "") or "").strip().lstrip("#")
            if font_color and re.fullmatch(r"[0-9A-Fa-f]{6}", font_color):
                font_color = f"FF{font_color}"
            inline_font = InlineFont(
                rFont=str(segment.get("font_name", "") or cell.font.name or "") or None,
                sz=float(str(segment.get("font_size", "") or cell.font.sz or 11).replace("pt", "").replace("px", "")),
                b=bool(segment.get("bold")) if "bold" in segment else bool(cell.font.bold),
                i=bool(segment.get("italic")) if "italic" in segment else bool(cell.font.italic),
                u="single" if bool(segment.get("underline")) else ("single" if cell.font.underline else None),
                strike=bool(segment.get("strike")) if "strike" in segment else bool(cell.font.strike),
                color=font_color or None,
                vertAlign="subscript" if vert_align == "subscript" else "superscript" if vert_align == "superscript" else None,
            )
            rich_text.append(TextBlock(inline_font, segment_text))
        return rich_text if wrote else None

    rich_value = excel_rich_text_value(style_payload.get("rich_runs"))
    cell.value = rich_value if rich_value is not None else (text if text != "" else None)
    if style_payload:
        _apply_excel_style(cell, style_payload)
        merge_range = str(style_payload.get("merge_range", "") or "").strip()
        if merge_range and re.fullmatch(r"[A-Z]+\d+:[A-Z]+\d+", merge_range):
            target_bounds = range_boundaries(merge_range)
            for merged_range in list(worksheet.merged_cells.ranges):
                existing_bounds = range_boundaries(str(merged_range))
                intersects = not (
                    existing_bounds[2] < target_bounds[0]
                    or target_bounds[2] < existing_bounds[0]
                    or existing_bounds[3] < target_bounds[1]
                    or target_bounds[3] < existing_bounds[1]
                )
                if intersects:
                    worksheet.unmerge_cells(str(merged_range))
            worksheet.merge_cells(merge_range)
            worksheet[merge_range.split(":", 1)[0]].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _apply_excel_page_setup(worksheet)
    workbook.save(path)
    workbook.close()


def update_word_template_item(path: Path, key: str, text: str, style_payload: dict[str, object]) -> None:
    from docx import Document  # noqa: E402

    document = Document(str(path))
    target_paragraph = None
    target_cell = None
    if key == "__append_paragraph__":
        target_paragraph = document.add_paragraph(text or "New paragraph")
    elif key.startswith("p:"):
        index = int(key.split(":", 1)[1])
        if index < 0 or index >= len(document.paragraphs):
            raise ValueError("Selected paragraph was not found in this format.")
        target_paragraph = document.paragraphs[index]
    elif key.startswith("t:"):
        match = re.fullmatch(r"t:(\d+):r:(\d+):c:(\d+)", key)
        if not match:
            raise ValueError("Selected table cell was not found in this format.")
        table_index, row_index, cell_index = [int(part) for part in match.groups()]
        try:
            cell = document.tables[table_index].rows[row_index].cells[cell_index]
        except IndexError as error:
            raise ValueError("Selected table cell was not found in this format.") from error
        target_cell = cell
        target_paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        if text != cell.text and not style_payload.get("paragraphs"):
            for extra in list(cell.paragraphs[1:]):
                extra._element.getparent().remove(extra._element)
    else:
        raise ValueError("Select a valid Word text block from the scanned format.")

    def set_paragraph_text_preserving_format(paragraph, value: str) -> None:
        if paragraph.runs:
            paragraph.runs[0].text = value
            for extra_run in paragraph.runs[1:]:
                extra_run.text = ""
        else:
            paragraph.add_run(value)

    def set_paragraph_rich_text(paragraph, runs_payload: object, fallback: str) -> bool:
        if not isinstance(runs_payload, list) or not runs_payload:
            return False
        base_run = paragraph.runs[0] if paragraph.runs else None
        base = {
            "bold": getattr(base_run.font, "bold", None) if base_run is not None else None,
            "italic": getattr(base_run.font, "italic", None) if base_run is not None else None,
            "underline": getattr(base_run.font, "underline", None) if base_run is not None else None,
            "strike": getattr(base_run.font, "strike", None) if base_run is not None else None,
            "name": getattr(base_run.font, "name", None) if base_run is not None else None,
            "size": getattr(base_run.font, "size", None) if base_run is not None else None,
            "color": _word_run_color_text(base_run) if base_run is not None else "",
        }
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        wrote_text = False
        for segment in runs_payload:
            if not isinstance(segment, dict):
                continue
            segment_text = str(segment.get("text", ""))
            if segment_text == "":
                continue
            wrote_text = True
            run = paragraph.add_run(segment_text)
            run.font.bold = bool(segment.get("bold")) if "bold" in segment else base["bold"]
            run.font.italic = bool(segment.get("italic")) if "italic" in segment else base["italic"]
            run.font.underline = bool(segment.get("underline")) if "underline" in segment else base["underline"]
            run.font.strike = bool(segment.get("strike")) if "strike" in segment else base["strike"]
            font_name = str(segment.get("font_name", "") or base["name"] or "").strip()
            if font_name:
                run.font.name = font_name
            font_size = segment.get("font_size", None)
            if font_size in (None, ""):
                if base["size"] is not None:
                    run.font.size = base["size"]
            else:
                from docx.shared import Pt  # noqa: E402

                try:
                    run.font.size = Pt(float(str(font_size).replace("pt", "").replace("px", "")))
                except Exception:
                    pass
            font_color = str(segment.get("font_color", "") or base["color"] or "").strip().lstrip("#")
            if font_color and re.fullmatch(r"[0-9A-Fa-f]{6}", font_color):
                from docx.shared import RGBColor  # noqa: E402

                run.font.color.rgb = RGBColor.from_string(font_color.upper())
            vert_align = str(segment.get("vert_align", "") or "").strip().lower()
            if vert_align == "subscript":
                run.font.subscript = True
                run.font.superscript = False
            elif vert_align == "superscript":
                run.font.superscript = True
                run.font.subscript = False
        if not wrote_text:
            paragraph.add_run(fallback)
        return True

    if target_cell is not None and isinstance(style_payload.get('paragraphs'), list):
        paragraphs = style_payload['paragraphs']
        for index, data in enumerate(paragraphs):
            p = target_cell.paragraphs[index] if index < len(target_cell.paragraphs) else target_cell.add_paragraph()
            value = str(data.get('text', ''))
            if p.text != value: set_paragraph_text_preserving_format(p, value)
        for extra in list(target_cell.paragraphs[len(paragraphs):]):
            extra._element.getparent().remove(extra._element)
        document.save(str(path))
        return
    rich_text_applied = set_paragraph_rich_text(target_paragraph, style_payload.get("rich_runs") if isinstance(style_payload, dict) else None, text)
    if not rich_text_applied:
        set_paragraph_text_preserving_format(target_paragraph, text)
    if style_payload:
        alignment = _normal_alignment_text(style_payload.get("horizontal", "") or style_payload.get("alignment", ""))
        if alignment:
            from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

            align_map = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            if alignment in align_map:
                target_paragraph.alignment = align_map[alignment]
        line_spacing = str(style_payload.get("line_spacing", "") or "").strip()
        if line_spacing:
            try:
                target_paragraph.paragraph_format.line_spacing = float(line_spacing)
            except Exception:
                pass
        indent = style_payload.get("indent", None)
        if indent is not None:
            from docx.shared import Pt  # noqa: E402

            try:
                target_paragraph.paragraph_format.left_indent = Pt(float(indent) * 12)
            except Exception:
                pass
        left_indent = str(style_payload.get("left_indent", "") or "").strip()
        first_line_indent = str(style_payload.get("first_line_indent", "") or "").strip()
        right_indent = str(style_payload.get("right_indent", "") or "").strip()
        if left_indent or first_line_indent or right_indent:
            from docx.shared import Pt  # noqa: E402

            def css_length_to_pt(raw: str) -> float | None:
                value = raw.strip().lower()
                try:
                    if value.endswith("px"):
                        return float(value[:-2]) * 0.75
                    if value.endswith("pt"):
                        return float(value[:-2])
                    if value:
                        return float(value) * 12
                except Exception:
                    return None
                return None

            left_pt = css_length_to_pt(left_indent)
            first_pt = css_length_to_pt(first_line_indent)
            right_pt = css_length_to_pt(right_indent)
            if left_pt is not None:
                target_paragraph.paragraph_format.left_indent = Pt(left_pt)
            if first_pt is not None:
                target_paragraph.paragraph_format.first_line_indent = Pt(first_pt)
            if right_pt is not None:
                target_paragraph.paragraph_format.right_indent = Pt(right_pt)
        left_tab = str(style_payload.get("left_tab", "") or "").strip()
        if left_tab:
            from docx.enum.text import WD_TAB_ALIGNMENT  # noqa: E402
            from docx.shared import Pt  # noqa: E402

            def css_length_to_pt(raw: str) -> float | None:
                value = raw.strip().lower()
                try:
                    if value.endswith("px"):
                        return float(value[:-2]) * 0.75
                    if value.endswith("pt"):
                        return float(value[:-2])
                    if value:
                        return float(value) * 12
                except Exception:
                    return None
                return None

            tab_pt = css_length_to_pt(left_tab)
            if tab_pt is not None and tab_pt >= 0:
                try:
                    target_paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(tab_pt), WD_TAB_ALIGNMENT.LEFT)
                except Exception:
                    pass
        for run in target_paragraph.runs:
            if "bold" in style_payload:
                run.font.bold = bool(style_payload.get("bold"))
            if "italic" in style_payload:
                run.font.italic = bool(style_payload.get("italic"))
            if "underline" in style_payload:
                run.font.underline = bool(style_payload.get("underline"))
            if "strike" in style_payload:
                run.font.strike = bool(style_payload.get("strike"))
            if "vert_align" in style_payload:
                run.font.subscript = str(style_payload.get("vert_align", "") or "") == "subscript"
                run.font.superscript = str(style_payload.get("vert_align", "") or "") == "superscript"
            font_name = str(style_payload.get("font_name", "") or "").strip()
            if font_name:
                run.font.name = font_name
            font_color = str(style_payload.get("font_color", "") or "").strip().lstrip("#")
            if font_color and re.fullmatch(r"[0-9A-Fa-f]{6}", font_color):
                from docx.shared import RGBColor  # noqa: E402

                run.font.color.rgb = RGBColor.from_string(font_color.upper())
            font_size = str(style_payload.get("font_size", "") or "").strip()
            if font_size:
                from docx.shared import Pt  # noqa: E402

                try:
                    run.font.size = Pt(float(font_size))
                except Exception:
                    pass
    document.save(str(path))


def update_template_item(
    kind: str,
    template_name: str,
    template_dir_text: str,
    user_id: int | None,
    key: str,
    text: str,
    style_payload: object,
) -> dict[str, object]:
    normalized_kind = kind.strip().lower()
    selected = _require_custom_template_entry(normalized_kind, template_name, template_dir_text, user_id)
    style = style_payload if isinstance(style_payload, dict) else {}
    working_path = editable_office_copy(selected.path)
    if normalized_kind == "statement":
        if working_path.suffix.lower() != ".xlsx":
            raise ValueError("Excel format editing needs an uploaded .xlsx file.")
        update_excel_template_item(working_path, key, text, style)
    elif normalized_kind == "certificate":
        if working_path.suffix.lower() != ".docx":
            raise ValueError("Word format editing needs an uploaded .docx file.")
        update_word_template_item(working_path, key, text, style)
    else:
        raise ValueError("Template kind must be statement or certificate.")
    meta = _rescan_and_save_template_meta(normalized_kind, selected.path)
    return {
        "saved": True,
        "kind": normalized_kind,
        "name": selected.name,
        "scan": meta.get("scan", {}),
    }


def update_template_batch(kind, name, template_dir, user_id, edits):
    if not isinstance(edits, list) or len(edits) > 10000:
        raise ValueError('Supply at most 10000 edited items.')
    selected = _require_custom_template_entry(kind, name, template_dir, user_id)
    original = editable_office_copy(selected.path)
    temporary = original.with_name('.edit-' + uuid4().hex + original.suffix)
    shutil.copy2(original, temporary)
    try:
        for edit in edits:
            if not isinstance(edit, dict) or not edit.get('key'):
                raise ValueError('Invalid format edit.')
            updater = update_excel_template_item if kind == 'statement' else update_word_template_item
            updater(temporary, str(edit['key']), str(edit.get('text', '')), dict(edit.get('style') or {}))
        temporary.replace(original)
        _rescan_and_save_template_meta(kind, selected.path)
    finally:
        temporary.unlink(missing_ok=True)
    return {'saved': True, 'name': name}


def create_statement_template_from_definition(name: str, definition: object, profile_payload: object, user_id: int | None = None) -> dict[str, object]:
    from openpyxl import Workbook  # noqa: E402
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
    from openpyxl.utils import get_column_letter  # noqa: E402

    base_name = safe_filename(" ".join(str(name).split()).strip() or "Custom Statement")
    profile = _coerce_template_profile(profile_payload)
    config = definition if isinstance(definition, dict) else {}
    headers_raw = config.get("headers", [])
    headers = [str(item).strip() for item in headers_raw if str(item).strip()] if isinstance(headers_raw, list) else []
    if not headers:
        headers = ["Date", "Description", "Cheque No.", "Debit", "Credit", "Balance"]
    max_cols = max(6, len(headers))
    transaction_rows = int(config.get("transaction_rows", 55) or 55)
    transaction_rows = min(1_000, max(10, transaction_rows))
    include_total = bool(config.get("include_total", False))
    include_summary = bool(config.get("include_summary", True))

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Statement"
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    worksheet.cell(1, 1, "Bank Statement").font = Font(bold=True, size=16)
    worksheet.cell(1, 1).alignment = Alignment(horizontal="center")
    top_labels = [
        ("Name", "customer_name"),
        ("Address", "customer_address"),
        ("A/c No.", "account_number"),
        ("Statement Period", "start_date"),
        ("Issue Date", "issue_date"),
    ]
    row_index = 3
    for label, key in top_labels:
        worksheet.cell(row_index, 1, label).font = Font(bold=True)
        worksheet.cell(row_index, 2, f"{{{{{key}}}}}")
        row_index += 1

    header_row = row_index + 1
    thin = Side(style="thin", color="FF9FB3C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill(fill_type="solid", fgColor="FFEAF4F6")
    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(header_row, column_index, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = fill
        cell.border = border
        worksheet.column_dimensions[get_column_letter(column_index)].width = 18 if column_index != 2 else 36

    for row in range(header_row + 1, header_row + transaction_rows + 1):
        for column_index in range(1, max_cols + 1):
            cell = worksheet.cell(row, column_index, "")
            cell.border = border
            if column_index in {1}:
                cell.number_format = "yyyy-mm-dd"
            elif column_index in {4, 5, 6}:
                cell.number_format = '#,##0.00'

    total_row = header_row + transaction_rows + 1
    if include_total:
        worksheet.cell(total_row, 1, "Total").font = Font(bold=True)
        for column_index in range(1, max_cols + 1):
            worksheet.cell(total_row, column_index).border = border
            worksheet.cell(total_row, column_index).font = Font(bold=True)
        if max_cols >= 6:
            worksheet.cell(total_row, 4, "{{total_debit}}")
            worksheet.cell(total_row, 5, "{{total_credit}}")
            worksheet.cell(total_row, 6, "{{closing_balance}}")

    if include_summary:
        summary_row = total_row + 3
        worksheet.cell(summary_row, 1, "Transaction Summary").font = Font(bold=True)
        worksheet.cell(summary_row + 1, 1, "Total Debit")
        worksheet.cell(summary_row + 1, 2, "{{total_debit}}")
        worksheet.cell(summary_row + 2, 1, "Total Credit")
        worksheet.cell(summary_row + 2, 2, "{{total_credit}}")
        worksheet.cell(summary_row + 3, 1, "Closing Balance")
        worksheet.cell(summary_row + 3, 2, "{{closing_balance}}")
        worksheet.cell(summary_row + 4, 1, "Balance In Words")
        worksheet.cell(summary_row + 4, 2, "{{balance_words}}")

    target_dir = _custom_template_dir("statement", user_id)
    output_path = target_dir / f"{base_name}.xlsx"
    workbook.save(output_path)
    workbook.close()
    meta = _rescan_and_save_template_meta("statement", output_path, profile)
    return {
        "kind": "statement",
        "name": base_name,
        "suffix": ".xlsx",
        "source": "custom",
        "editable": True,
        "scan_summary": dict(meta.get("scan", {})).get("summary", {}) if isinstance(meta.get("scan"), dict) else {},
    }


def _normalize_description_mode(value: str) -> str:
    value = value.strip()
    if value in DESCRIPTION_MODE_OPTIONS:
        return DESCRIPTION_MODE_OPTIONS[value]
    if value in DESCRIPTION_MODE_LABELS:
        return value
    return DESCRIPTION_MODE_OPTIONS["Label + Name"]


def _parse_optional_date(value: str) -> date | None:
    value = value.strip()
    if not value:
        return None
    try:
        return parse_iso_date(value)
    except ValueError as error:
        raise ValueError("Invalid opening date format. Expected YYYY-MM-DD.") from error


def _safe_float(value: str, default: float = 0.0) -> float:
    cleaned = value.strip().replace(",", "")
    if not cleaned:
        return default
    try:
        return float(cleaned)
    except ValueError:
        return default


def _safe_int(value: str, default: int = 0) -> int:
    cleaned = value.strip()
    if not cleaned:
        return default
    try:
        return int(float(cleaned))
    except ValueError:
        return default


def _safe_iso_date(value: str, field_name: str) -> date:
    cleaned = value.strip()
    if not cleaned:
        raise ValueError(f"{field_name} is required.")
    try:
        return parse_iso_date(cleaned)
    except ValueError as error:
        raise ValueError(f"Invalid {field_name} format. Expected YYYY-MM-DD.") from error


def _parse_yes_no(value: str, default: bool = True) -> bool:
    normalized = value.strip().lower()
    if not normalized:
        return default
    return normalized not in {"no", "false", "0", "off"}


def _normalize_amount_rounding_mode(value: str) -> str:
    normalized = str(value or "").strip()
    if normalized in AMOUNT_ROUNDING_OPTIONS:
        return AMOUNT_ROUNDING_OPTIONS[normalized]
    if normalized in AMOUNT_ROUNDING_OPTIONS.values():
        return normalized
    return AMOUNT_ROUNDING_OPTIONS["Automatic (Present Rule)"]


def _parse_amount_limit(value: str, default: int) -> int:
    cleaned = str(value or "").replace(",", "").strip()
    if not cleaned:
        return default
    try:
        return max(1, int(round(float(cleaned))))
    except ValueError:
        return default


def _parse_monthly_transaction_counts(value: object) -> dict[tuple[int, int], int]:
    raw = value
    if isinstance(raw, str):
        try:
            raw = json.loads(raw) if raw.strip() else {}
        except Exception:
            raw = {}
    if not isinstance(raw, dict):
        return {}
    counts: dict[tuple[int, int], int] = {}
    for key, amount in raw.items():
        try:
            year_text, month_text = str(key).split("-", 1)
            year = int(year_text)
            month = int(month_text)
            count = int(float(str(amount)))
        except Exception:
            continue
        if 1900 <= year <= 2200 and 1 <= month <= 12 and count > 0:
            counts[(year, month)] = max(0, min(50, count))
    return counts


def build_config(form_data: dict[str, object], forced_seed: int | None = None, user_id: int | None = None) -> StatementConfig:
    defaults = default_form_values()

    def text(key: str) -> str:
        raw = form_data.get(key, defaults.get(key, ""))
        return str(raw).strip()

    rules = load_persistent_rules(user_id)
    prepend_mode = _parse_yes_no(text("prepend_statement_mode"), default=False)
    prepend_start_text = text("prepend_start_date")
    prepend_anchor_text = text("prepend_anchor_date")
    start_text = prepend_start_text if prepend_mode and prepend_start_text else text("start_date")
    end_text = text("end_date")
    seed_text = str(forced_seed) if forced_seed is not None else text("seed")

    start_date = _safe_iso_date(start_text, "Start date")
    end_date = _safe_iso_date(end_text, "End date")
    prepend_start_date = _safe_iso_date(prepend_start_text, "Previous statement start date") if prepend_mode and prepend_start_text else None
    prepend_anchor_date = _safe_iso_date(prepend_anchor_text, "Existing statement start date") if prepend_mode and prepend_anchor_text else None

    return StatementConfig(
        bank_name=text("bank_name"),
        branch_name=text("branch_name"),
        customer_name=text("customer_name"),
        customer_address=text("customer_address"),
        account_number=text("account_number"),
        account_type=text("account_type"),
        member_id=text("member_id"),
        currency=text("currency") or "NPR",
        reference_no=text("reference_no"),
        opening_date=_parse_optional_date(text("opening_date")),
        start_date=start_date,
        end_date=end_date,
        opening_balance=_safe_float(text("opening_balance")),
        target_closing_balance=_safe_float(text("target_closing_balance")),
        prepend_statement_mode=prepend_mode,
        prepend_start_date=prepend_start_date,
        prepend_anchor_date=prepend_anchor_date,
        prepend_anchor_balance=_safe_float(text("prepend_anchor_balance")) if prepend_mode else 0.0,
        deposit_min_amount=_parse_amount_limit(text("deposit_min_amount"), 15_000),
        deposit_max_amount=_parse_amount_limit(text("deposit_max_amount"), 99_000),
        withdrawal_min_amount=_parse_amount_limit(text("withdrawal_min_amount"), 15_000),
        withdrawal_max_amount=_parse_amount_limit(text("withdrawal_max_amount"), 65_000),
        amount_rounding_mode=_normalize_amount_rounding_mode(text("amount_rounding_mode")),
        amount_rounding_percentages=(parse_percentages(form_data.get("amount_rounding_percentages", ""))
                                     if _normalize_amount_rounding_mode(text("amount_rounding_mode")) == "custom" else {}),
        interest_rate=_safe_float(text("interest_rate")),
        tax_rate=_safe_float(text("tax_rate")),
        cheque_start=_safe_int(text("cheque_start")),
        include_cheque_column=_parse_yes_no(text("include_cheque_column"), default=True),
        deposit_text=text("deposit_text"),
        withdrawal_text=text("withdrawal_text"),
        description_extra_text=text("description_extra_text"),
        interest_text=text("interest_text"),
        tax_text=text("tax_text"),
        date_column_mode=text("date_column_mode").lower() if text("date_column_mode").lower() in {"single", "txn_value"} else "single",
        first_date_description=text("first_date_description"),
        last_date_description=text("last_date_description"),
        closing_row_mode=text("closing_row_mode").lower() if text("closing_row_mode").lower() in {"transaction_allowed", "description_only"} else "description_only",
        statement_row_mode=text("statement_row_mode").lower() if text("statement_row_mode").lower() in {"auto", "custom"} else "auto",
        statement_row_count=_safe_int(text("statement_row_count")) if text("statement_row_mode").lower() == "custom" and text("statement_row_count") else None,
        transaction_count_mode=text("transaction_count_mode").lower() if text("transaction_count_mode").lower() in {"auto", "custom"} else "auto",
        monthly_transaction_counts=_parse_monthly_transaction_counts(form_data.get("monthly_transaction_counts", "")),
        deposit_names=names_from_text(text("deposit_names")),
        withdrawal_names=names_from_text(text("withdrawal_names")),
        deposit_name_mode=_normalize_description_mode(text("deposit_mode")),
        withdrawal_name_mode=_normalize_description_mode(text("withdrawal_mode")),
        holiday_dates=blocked_dates(start_text, end_text, rules["custom_holidays"], rules["excluded_saturdays"]),
        quarter_date_overrides=_quarter_override_lookup(merged_quarter_override_map(start_text, end_text, user_id)),
        seed=_safe_int(seed_text) if seed_text else None,
    )


def serialize_result(config: StatementConfig, result) -> dict[str, object]:
    gap = result.final_balance - config.target_closing_balance
    seed_label = str(getattr(result, "seed_label", "") or result.seed)
    rows = []
    for row in result.rows:
        rows.append(
            {
                "date": row.date.isoformat(),
                "txn_date": row.date.isoformat(),
                "value_date": row.date.isoformat(),
                "description": row.description,
                "cheque_no": row.cheque_no,
                "debit": row.debit,
                "credit": row.credit,
                "balance": row.balance,
                "debit_text": format_amount(row.debit) if row.debit else "",
                "credit_text": format_amount(row.credit) if row.credit else "",
                "balance_text": format_amount(row.balance),
                "category": row.category,
                "is_system": row.is_system,
            }
        )
    return {
        "generated_seed": result.seed,
        "issue_date": result.issue_date.isoformat(),
        "last_transaction_date": result.last_transaction_date.isoformat(),
        "opening_business_date": result.opening_business_date.isoformat(),
        "ending_business_date": result.ending_business_date.isoformat(),
        "final_balance": result.final_balance,
        "is_manual_edit": bool(getattr(result, "is_manual_edit", False)),
        "summary": {
            "final_balance_text": f"Rs. {format_amount(result.final_balance)}",
            "issue_date": result.issue_date.isoformat(),
            "row_count": len(result.rows),
            "target_gap_text": f"Rs. {format_amount(gap)}",
            "total_deposits_text": f"Rs. {format_amount(result.summary.total_deposits)}",
            "total_withdrawals_text": f"Rs. {format_amount(result.summary.total_withdrawals)}",
            "total_interest_text": f"Rs. {format_amount(result.summary.total_interest)}",
            "total_tax_text": f"Rs. {format_amount(result.summary.total_tax)}",
            "seed_used": seed_label,
        },
        "rows": rows,
    }


def serialize_preview_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    preview_rows: list[dict[str, object]] = []
    for row in rows:
        debit = round_money(float(row.get("debit", 0) or 0))
        credit = round_money(float(row.get("credit", 0) or 0))
        balance = round_money(float(row.get("balance", 0) or 0))
        preview_rows.append(
            {
                "date": str(row.get("date", "")),
                "txn_date": str(row.get("txn_date", row.get("date", ""))),
                "value_date": str(row.get("value_date", row.get("date", ""))),
                "description": str(row.get("description", "")),
                "cheque_no": str(row.get("cheque_no", "")),
                "debit": debit,
                "credit": credit,
                "balance": balance,
                "debit_text": format_amount(debit) if debit else "",
                "credit_text": format_amount(credit) if credit else "",
                "balance_text": format_amount(balance),
                "category": str(row.get("category", "")),
                "is_system": bool(row.get("is_system", False)),
            }
        )
    return preview_rows


def hydrate_result_payload(config: StatementConfig, result_payload: dict[str, object]):
    errors = validate_transaction_dates(config, result_payload.get("rows", []))
    if errors:
        raise ValueError(str(errors[0]["message"]) + " Recalculate the statement before exporting.")
    opening_business_date = parse_iso_date(str(result_payload.get("opening_business_date") or resolve_business_day(config.start_date, config.holiday_dates).isoformat()))
    ending_business_date = parse_iso_date(str(result_payload.get("ending_business_date") or resolve_business_day(config.end_date, config.holiday_dates).isoformat()))

    from statement_generator.generator import StatementResult, StatementRow, StatementSummary  # noqa: E402

    rows: list[StatementRow] = []
    summary = StatementSummary()
    last_balance = config.opening_balance
    for row in result_payload.get("rows", []):
        if not isinstance(row, dict) or not str(row.get("date", "")).strip():
            continue
        statement_row = StatementRow(
            date=parse_iso_date(str(row.get("date", ""))),
            description=str(row.get("description", "")),
            cheque_no=str(row.get("cheque_no", "")),
            debit=float(row.get("debit", 0) or 0),
            credit=float(row.get("credit", 0) or 0),
            balance=float(row.get("balance", 0) or 0),
            category=str(row.get("category", "")),
            is_system=bool(row.get("is_system", False)),
        )
        rows.append(statement_row)
        last_balance = statement_row.balance
        if statement_row.category == "deposit":
            summary.total_deposits = round_money(summary.total_deposits + statement_row.credit)
            summary.deposit_count += 1
        elif statement_row.category == "withdrawal":
            summary.total_withdrawals = round_money(summary.total_withdrawals + statement_row.debit)
            summary.withdrawal_count += 1
        elif statement_row.category == "interest":
            summary.total_interest = round_money(summary.total_interest + statement_row.credit)
        elif statement_row.category == "tax":
            summary.total_tax = round_money(summary.total_tax + statement_row.debit)

    final_balance = float(result_payload.get("final_balance", last_balance) or last_balance)
    generated_seed_raw = result_payload.get("generated_seed", 0)
    seed_label = str(dict(result_payload.get("summary", {})).get("seed_used") or generated_seed_raw)
    return StatementResult(
        rows=rows,
        summary=summary,
        events=[],
        final_balance=final_balance,
        opening_business_date=opening_business_date,
        ending_business_date=ending_business_date,
        last_transaction_date=parse_iso_date(str(result_payload.get("last_transaction_date") or ending_business_date.isoformat())),
        issue_date=parse_iso_date(str(result_payload.get("issue_date") or next_business_day(ending_business_date, config.holiday_dates, include_self=False).isoformat())),
        seed=int(generated_seed_raw) if str(generated_seed_raw).strip().isdigit() else 0,
        seed_label=seed_label,
        is_manual_edit=bool(result_payload.get("is_manual_edit", False)),
    )


def build_export_input(payload: dict[str, object], user_id: int | None = None):
    form = payload.get("config")
    if not isinstance(form, dict):
        raise ValueError("Export request is missing config data.")
    generated_seed = payload.get("generated_seed")
    forced_seed = int(str(generated_seed)) if str(generated_seed).strip() else None
    config = build_config(form, forced_seed=forced_seed, user_id=user_id)
    result_payload = payload.get("result_payload")
    if isinstance(result_payload, dict):
        result = hydrate_result_payload(config, result_payload)
    else:
        result = generate_statement(config)

    rate_mode_label = str(payload.get("rate_mode", form.get("rate_mode", default_form_values()["rate_mode"]))).strip()
    rate_mode = "manual" if rate_mode_label == "Manual" else "auto"
    manual_rate_text = str(payload.get("manual_rate", form.get("manual_rate", ""))).strip()
    manual_rate = float(manual_rate_text.replace(",", "")) if manual_rate_text else None
    rate_type = str(payload.get("rate_type", form.get("rate_type", "sell"))).strip() or "sell"
    try:
        exchange_rate = resolve_exchange_rate(result.issue_date, mode=rate_mode, manual_rate=manual_rate, rate_type=rate_type)
    except Exception as error:
        if rate_mode == "auto" and manual_rate and manual_rate > 0 and _is_exchange_rate_error(str(error)):
            exchange_rate = resolve_exchange_rate(result.issue_date, mode="manual", manual_rate=manual_rate, rate_type=rate_type)
        else:
            raise
    export_payload = build_payload(config, result, exchange_rate)
    return config, result, exchange_rate, export_payload


def json_bytes(payload: dict[str, object]) -> bytes:
    return json.dumps(payload, indent=2, default=json_default).encode("utf-8")


def build_normal_statement_bytes(export_payload: dict[str, object]) -> bytes:
    account = export_payload["account"]
    statement = export_payload["statement"]
    include_cheque = bool(statement.get("include_cheque_column", True))
    two_date_columns = str(statement.get("date_column_mode", "single")) == "txn_value"
    date_colspan = 2 if two_date_columns else 1
    table_colspan = date_colspan + 4 + (1 if include_cheque else 0)
    rows_html = []
    for row in export_payload["statement_rows"]:
        debit_text = str(row.get("debit_text") or format_amount(float(row["debit"]))) if float(row["debit"]) > 0 else ""
        credit_text = str(row.get("credit_text") or format_amount(float(row["credit"]))) if float(row["credit"]) > 0 else ""
        balance_text = str(row.get("balance_text") or format_amount(float(row["balance"])))
        date_cells = (
            f"<td class='date-cell' style='mso-number-format:yyyy\\-mm\\-dd;white-space:nowrap'>{escape(str(row.get('txn_date') or row['date']))}</td>"
            f"<td class='date-cell' style='mso-number-format:yyyy\\-mm\\-dd;white-space:nowrap'>{escape(str(row.get('value_date') or row['date']))}</td>"
            if two_date_columns
            else f"<td class='date-cell' style='mso-number-format:yyyy\\-mm\\-dd;white-space:nowrap'>{escape(str(row['date']))}</td>"
        )
        cheque_cell = (
            f"<td style='mso-number-format:0'>{escape(str(row['cheque_no']))}</td>"
            if include_cheque
            else ""
        )
        rows_html.append(
            "<tr>"
            f"{date_cells}"
            f"<td style=\"mso-number-format:'\\@'\">{escape(str(row['description']))}</td>"
            f"{cheque_cell}"
            f"<td style=\"text-align:right;mso-number-format:'_-* #,##0.00_-;\\\\-* #,##0.00_-;_-* &quot;-&quot;??_-;_-@_-'\">{escape(debit_text)}</td>"
            f"<td style=\"text-align:right;mso-number-format:'_-* #,##0.00_-;\\\\-* #,##0.00_-;_-* &quot;-&quot;??_-;_-@_-'\">{escape(credit_text)}</td>"
            f"<td style=\"text-align:right;mso-number-format:'_-* #,##0.00_-;\\\\-* #,##0.00_-;_-* &quot;-&quot;??_-;_-@_-'\">{escape(balance_text)}</td>"
            "</tr>"
        )
    date_headers = "<th class='date-cell'>TXN Date</th><th class='date-cell'>Value Date</th>" if two_date_columns else "<th class='date-cell'>Date</th>"
    cheque_header = "<th>Cheque No.</th>" if include_cheque else ""
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@page {{ size: A4; margin: 10mm; }}
body {{ font-family: Calibri, Arial, sans-serif; color: #16324a; margin: 0; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 1px solid #a9b7c4; padding: 8px; }}
th {{ background: #d9eaf7; text-transform: uppercase; font-size: 11px; }}
.date-cell {{ white-space: nowrap; }}
.title {{ background: #1f4e78; color: white; font-size: 18px; font-weight: 700; text-align: center; }}
.meta td {{ border: none; padding: 5px 2px; }}
</style>
</head>
<body>
<table>
  <tr><td class="title" colspan="{table_colspan}">BANK STATEMENT</td></tr>
</table>
<table class="meta" style="margin-top: 12px;">
  <tr><td><strong>Name</strong></td><td>{escape(str(account['customer_name']))}</td><td><strong>Account No.</strong></td><td style='mso-number-format:0'>{escape(str(account['account_number']))}</td></tr>
  <tr><td><strong>Address</strong></td><td>{escape(str(account['customer_address']))}</td><td><strong>Account Type</strong></td><td>{escape(str(account['account_type']))}</td></tr>
  <tr><td><strong>Period</strong></td><td>{escape(str(statement['period_label_iso']))}</td><td><strong>Issue Date</strong></td><td>{escape(str(statement['issue_date_iso']))}</td></tr>
  <tr><td><strong>Final Balance</strong></td><td colspan="3">{escape(str(export_payload["summary"]['final_balance_text']))}</td></tr>
</table>
<table style="margin-top: 16px;">
  <thead>
    <tr>
      {date_headers}
      <th>Description</th>
      {cheque_header}
      <th>Debit</th>
      <th>Credit</th>
      <th>Balance</th>
    </tr>
  </thead>
  <tbody>
    {''.join(rows_html)}
  </tbody>
</table>
</body>
</html>"""
    return sample_html(html).encode("utf-8")


def build_normal_certificate_bytes(export_payload: dict[str, object]) -> bytes:
    account = export_payload["account"]
    statement = export_payload["statement"]
    rates = export_payload["rates"]
    certificate = export_payload["certificate"]
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@page {{ size: A4; margin: 12mm; }}
body {{ font-family: Cambria, Georgia, serif; color: #16324a; line-height: 1.65; margin: 0; }}
h1 {{ text-align: center; letter-spacing: 0.05em; }}
.meta {{ margin-top: 24px; }}
.meta td {{ padding: 6px 10px 6px 0; vertical-align: top; }}
</style>
</head>
<body>
  <h1>BALANCE CERTIFICATE</h1>
  <p>This is to certify that <strong>{escape(str(account['customer_name']))}</strong>, address <strong>{escape(str(account['customer_address']))}</strong>, has maintained account number <strong>{escape(str(account['account_number']))}</strong> with <strong>{escape(str(account['bank_name']))}</strong>, {escape(str(account['branch_name']))}.</p>
  <p>As of <strong>{escape(str(statement['as_of_long']))}</strong>, the available balance in the account is <strong>NPR {escape(str(certificate['total_balance_npr_text']))}</strong> ({escape(str(certificate['balance_words_npr']))}).</p>
  <p>Using the issue-day USD/NPR exchange rate of <strong>{escape(str(rates['usd_npr_text']))}</strong>, the equivalent balance is <strong>USD {escape(str(certificate['equivalent_usd_text']))}</strong> ({escape(str(certificate['balance_words_usd']))}).</p>
  <table class="meta">
    <tr><td><strong>Issue Date</strong></td><td>{escape(str(statement['issue_date_long']))}</td></tr>
    <tr><td><strong>Account Type</strong></td><td>{escape(str(account['account_type']))}</td></tr>
    <tr><td><strong>Reference No.</strong></td><td>{escape(str(account['reference_no']))}</td></tr>
    <tr><td><strong>Rate Source</strong></td><td>{escape(str(rates['source_label']))}</td></tr>
  </table>
  <p>This certificate is issued upon the request of the account holder for record purposes.</p>
</body>
</html>"""
    return sample_html(html).encode("utf-8")


def _openpyxl_color(value) -> str:
    try:
        rgb = str(value.rgb or "")
    except Exception:
        return ""
    if len(rgb) == 8:
        rgb = rgb[2:]
    return f"#{rgb}" if re.fullmatch(r"[0-9A-Fa-f]{6}", rgb) else ""


def _openpyxl_border_side_style(side) -> str:
    try:
        style = str(side.style or "")
    except Exception:
        return ""
    if not style:
        return ""
    width = "2px" if style in {"medium", "thick", "double"} else "1px"
    color = _openpyxl_color(getattr(side, "color", None)) or "#16324a"
    return f"{width} solid {color}"


def _excel_cell_number_format(cell) -> str:
    try:
        return str(cell.number_format or "")
    except Exception:
        return ""


def _excel_cell_has_date_format(cell) -> bool:
    try:
        if bool(getattr(cell, "is_date", False)):
            return True
    except Exception:
        pass
    number_format = _excel_cell_number_format(cell).lower()
    return bool(number_format and re.search(r"(^|[^a-z])[dmy]+([^a-z]|$)", number_format))


def _excel_cell_has_money_format(cell) -> bool:
    number_format = _excel_cell_number_format(cell)
    return bool(number_format and ("#,##" in number_format or "0.00" in number_format))


def _excel_preview_cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and _excel_cell_has_money_format(cell):
        return f"{float(value):,.2f}"
    return str(value)


def _excel_preview_cell_style(cell) -> str:
    styles: list[str] = []
    font = cell.font
    fill = cell.fill
    alignment = cell.alignment
    border = cell.border
    if getattr(font, "name", None):
        styles.append(f"font-family:{font.name}")
    if getattr(font, "sz", None):
        styles.append(f"font-size:{font.sz}pt")
    if getattr(font, "bold", False):
        styles.append("font-weight:700")
    if getattr(font, "italic", False):
        styles.append("font-style:italic")
    if getattr(font, "underline", None):
        styles.append("text-decoration:underline")
    font_color = _openpyxl_color(getattr(font, "color", None))
    if font_color:
        styles.append(f"color:{font_color}")
    fill_color = _openpyxl_color(getattr(fill, "fgColor", None))
    if fill_color and fill_color.lower() != "#000000" and getattr(fill, "fill_type", None):
        styles.append(f"background-color:{fill_color}")
    if getattr(alignment, "horizontal", None):
        styles.append(f"text-align:{alignment.horizontal}")
    if getattr(alignment, "vertical", None):
        vertical = "middle" if alignment.vertical == "center" else alignment.vertical
        styles.append(f"vertical-align:{vertical}")
    if getattr(alignment, "wrap_text", False):
        styles.append("white-space:pre-wrap")
    if _excel_cell_has_date_format(cell):
        styles.append("white-space:nowrap")
    try:
        side_map = {
            "top": border.top,
            "right": border.right,
            "bottom": border.bottom,
            "left": border.left,
        }
        border_sides = {
            side_name: css
            for side_name, side in side_map.items()
            if (css := _openpyxl_border_side_style(side))
        }
        if border_sides and len(border_sides) == 4 and len(set(border_sides.values())) == 1:
            styles.append(f"border:{next(iter(border_sides.values()))}")
        else:
            for side_name, css in border_sides.items():
                styles.append(f"border-{side_name}:{css}")
    except Exception:
        pass
    return ";".join(styles)


def build_excel_preview_html_bytes(path: Path, title: str = "Statement") -> bytes:
    from openpyxl import load_workbook  # noqa: E402
    from openpyxl.utils import get_column_letter, range_boundaries  # noqa: E402

    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    meaningful_max_row, meaningful_max_col = _worksheet_meaningful_bounds(worksheet)
    merged_starts: dict[str, tuple[int, int]] = {}
    covered: set[str] = set()
    for merged_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if min_row > meaningful_max_row or min_col > meaningful_max_col:
            continue
        max_row = min(max_row, meaningful_max_row)
        max_col = min(max_col, meaningful_max_col)
        start = f"{get_column_letter(min_col)}{min_row}"
        merged_starts[start] = (max_col - min_col + 1, max_row - min_row + 1)
        for row_index in range(min_row, max_row + 1):
            for col_index in range(min_col, max_col + 1):
                address = f"{get_column_letter(col_index)}{row_index}"
                if address != start:
                    covered.add(address)

    max_row = meaningful_max_row
    max_col = min(meaningful_max_col, 80)
    colgroup = []
    for col_index in range(1, max_col + 1):
        letter = get_column_letter(col_index)
        width = worksheet.column_dimensions[letter].width
        pixel_width = max(36, round(float(width or 10) * 7 + 8))
        colgroup.append(f'<col style="width:{pixel_width}px">')

    rows_html = []
    for row_index in range(1, max_row + 1):
        height = worksheet.row_dimensions[row_index].height
        row_style = f' style="height:{round(float(height) * 4 / 3)}px"' if height else ""
        cells_html = []
        for col_index in range(1, max_col + 1):
            cell = worksheet.cell(row=row_index, column=col_index)
            if cell.coordinate in covered:
                continue
            colspan, rowspan = merged_starts.get(cell.coordinate, (1, 1))
            attrs = []
            if colspan > 1:
                attrs.append(f'colspan="{colspan}"')
            if rowspan > 1:
                attrs.append(f'rowspan="{rowspan}"')
            style = _excel_preview_cell_style(cell)
            if style:
                attrs.append(f'style="{escape(style, quote=True)}"')
            value = _excel_preview_cell_text(cell)
            cells_html.append(f"<td {' '.join(attrs)}>{escape(value)}</td>")
        rows_html.append(f"<tr{row_style}>{''.join(cells_html)}</tr>")
    workbook.close()
    html = f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>{escape(title)}</title>
<style>
@page {{ size: A4; margin: 10mm; }}
body {{ margin: 0; background: #fff; font-family: Calibri, Arial, sans-serif; color: #16324a; }}
table {{ border-collapse: collapse; table-layout: fixed; width: 100%; }}
td {{ min-width: 28px; padding: 4px 6px; white-space: pre-wrap; }}
</style>
</head>
<body><table><colgroup>{''.join(colgroup)}</colgroup><tbody>{''.join(rows_html)}</tbody></table></body>
</html>"""
    return sample_html(html).encode("utf-8")


def build_docx_preview_html_bytes(path: Path, title: str = "Balance Certificate") -> bytes:
    return sample_html(word_html(editable_office_copy(path), title)).encode("utf-8")


def extract_bearer_token(headers) -> str:
    raw = str(headers.get("Authorization", "")).strip()
    if raw.lower().startswith("bearer "):
        return raw[7:].strip()
    return ""


def _read_login_attempts() -> dict[str, object]:
    if not LOGIN_ATTEMPT_FILE.exists():
        return {}
    try:
        payload = json.loads(LOGIN_ATTEMPT_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def _write_login_attempts(attempts: dict[str, object]) -> None:
    LOGIN_ATTEMPT_FILE.write_text(json.dumps(attempts, indent=2), encoding="utf-8")


def _login_attempt_key(client_ip: str, username: str) -> str:
    normalized_user = username.strip().lower()
    return hashlib.sha256(f"{client_ip}|{normalized_user}".encode("utf-8")).hexdigest()


def assert_login_not_rate_limited(client_ip: str, username: str) -> None:
    with LOGIN_ATTEMPT_LOCK:
        attempts = _read_login_attempts()
        now = int(time.time())
        key = _login_attempt_key(client_ip, username)
        record = attempts.get(key) if isinstance(attempts.get(key), dict) else {}
        locked_until = int(record.get("locked_until", 0) or 0)
        if locked_until > now:
            minutes = max(1, -(-(locked_until - now) // 60))
            raise PermissionError(
                f"Too many failed login attempts. Try again after {minutes} minute(s)."
            )
        failures = [
            int(ts)
            for ts in record.get("failures", [])
            if str(ts).lstrip("-").isdigit() and int(ts) >= now - LOGIN_WINDOW_SECONDS
        ]
        attempts[key] = {"failures": failures, "locked_until": 0}
        _write_login_attempts(attempts)


def record_login_attempt(client_ip: str, username: str, success: bool) -> None:
    with LOGIN_ATTEMPT_LOCK:
        attempts = _read_login_attempts()
        key = _login_attempt_key(client_ip, username)
        if success:
            attempts.pop(key, None)
            _write_login_attempts(attempts)
            return
        now = int(time.time())
        record = attempts.get(key) if isinstance(attempts.get(key), dict) else {}
        failures = [
            int(ts)
            for ts in record.get("failures", [])
            if str(ts).lstrip("-").isdigit() and int(ts) >= now - LOGIN_WINDOW_SECONDS
        ]
        failures.append(now)
        locked_until = now + LOGIN_WINDOW_SECONDS if len(failures) >= LOGIN_MAX_FAILURES else 0
        attempts[key] = {"failures": failures, "locked_until": locked_until}
        _write_login_attempts(attempts)


def build_user_workspace_payload(current_user: AuthUser, start_date: str, end_date: str) -> dict[str, object]:
    return {
        "current_user": current_user.as_payload(),
        "history": auth_store().statement_history(current_user),
        "users": auth_store().list_users() if current_user.is_admin else [],
        "devices": auth_store().devices_for_user(current_user, current_user.id),
        "holidays": build_holiday_payload(start_date, end_date, view="All", user_id=current_user.id),
        "posting_dates": build_posting_date_payload(start_date, end_date, user_id=current_user.id),
        "activities": auth_store().recent_activities(current_user) if current_user.is_admin else {"activities": []},
    }


def _is_exchange_rate_error(message: str) -> bool:
    lower = message.lower()
    return (
        "usd/npr exchange rate" in lower
        or "exchange rate service" in lower
        or "nepal rastra bank" in lower
        or "could not fetch exchange rate" in lower
    )


class StatementWebHandler(BaseHTTPRequestHandler):
    server_version = "StatementWebGenerator/1.0"

    def log_message(self, format: str, *args) -> None:  # noqa: A003
        print(f"[{self.log_date_time_string()}] {self.address_string()} - {format % args}")

    def do_GET(self) -> None:  # noqa: N802
        try:
            parsed = urlparse(self.path)
            route = parsed.path
            if route in {"/", "/index.html"}:
                self._serve_file(APP_ROOT / "templates" / "index.html", "text/html; charset=utf-8")
                return
            if route.startswith("/static/"):
                relative = route.removeprefix("/static/")
                self._serve_file(APP_ROOT / "static" / relative, base=APP_ROOT / "static")
                return
            if route.startswith("/site_integration/"):
                relative = route.removeprefix("/site_integration/")
                self._serve_file(APP_ROOT / "site_integration" / relative, base=APP_ROOT / "site_integration")
                return
            if route == "/api/letterhead":
                current_user = self._require_statement_user()
                query = parse_qs(parsed.query)
                value = load_letterhead(APP_ROOT / 'letterheads', current_user.id, query.get('kind', [''])[0], query.get('name', [''])[0])
                self._send_json({'letterhead': value})
                return
            if route == "/api/bootstrap":
                current_user = self._require_user()
                defaults = bootstrap_defaults(current_user)
                try:
                    live_rate = resolve_exchange_rate(date.today(), mode="auto", manual_rate=None, rate_type="sell", timeout=5)
                except Exception:
                    live_rate = None
                payload = {
                    "current_user": current_user.as_payload(),
                    "defaults": defaults,
                    "description_modes": list(DESCRIPTION_MODE_OPTIONS.keys()),
                    "amount_rounding_modes": AMOUNT_ROUNDING_OPTIONS,
                    "rate_modes": ["Auto (NRB)", "Manual"],
                    "rate_types": ["sell", "buy"],
                    "rate_sources": {
                        "page": NRB_FOREX_PAGE,
                        "docs": NRB_FOREX_DOCS,
                    },
                    "live_rate": live_rate,
                    "templates": serialize_catalog(defaults["template_dir"], current_user.id),
                    "profile_formats": auth_store().list_profile_formats(current_user),
                    **build_user_workspace_payload(current_user, defaults["start_date"], defaults["end_date"]),
                }
                self._send_json(payload)
                return
            if route == "/api/auto_refresh":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                start_date = (query.get("start_date") or [default_form_values()["start_date"]])[0]
                end_date = (query.get("end_date") or [default_form_values()["end_date"]])[0]
                refresh = refresh_from_internet(start_date, end_date, current_user.id)
                try:
                    live_rate = resolve_exchange_rate(date.today(), mode="auto", manual_rate=None, rate_type="sell", timeout=5)
                except Exception:
                    live_rate = None
                self._send_json(
                    {
                        "current_user": current_user.as_payload(),
                        "holidays": build_holiday_payload(start_date, end_date, view="All", user_id=current_user.id),
                        "posting_dates": build_posting_date_payload(start_date, end_date, user_id=current_user.id),
                        "live_rate": live_rate,
                        "refresh": refresh,
                    }
                )
                return
            if route == "/api/profile":
                current_user = self._require_user()
                self._send_json(auth_store().load_profile(current_user))
                return
            if route == "/api/profile_formats":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                name = (query.get("name") or [""])[0].strip()
                if name:
                    self._send_json(auth_store().load_profile_format(current_user, name))
                else:
                    self._send_json(auth_store().list_profile_formats(current_user))
                return
            if route == "/api/history":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                selected_user_text = (query.get("user_id") or [""])[0]
                selected_user_id = int(selected_user_text) if selected_user_text.strip() else None
                self._send_json(auth_store().statement_history(current_user, selected_user_id))
                return
            if route == "/api/statement_detail":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                statement_id = int((query.get("id") or ["0"])[0] or 0)
                self._send_json(auth_store().statement_detail(current_user, statement_id))
                return
            if route == "/api/users":
                current_user = self._require_user()
                if not current_user.is_admin:
                    raise PermissionError("Only the admin account can view the user list.")
                self._send_json({"users": auth_store().list_users()})
                return
            if route == "/api/activities":
                current_user = self._require_admin()
                query = parse_qs(parsed.query)
                selected_user_text = (query.get("user_id") or [""])[0]
                selected_user_id = int(selected_user_text) if selected_user_text.strip() else None
                self._send_json(auth_store().recent_activities(current_user, selected_user_id))
                return
            if route == "/api/devices":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                selected_user_text = (query.get("user_id") or [""])[0]
                selected_user_id = int(selected_user_text) if selected_user_text.strip() else None
                self._send_json(auth_store().devices_for_user(current_user, selected_user_id))
                return
            if route == "/api/holidays":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                start_date = (query.get("start_date") or [default_form_values()["start_date"]])[0]
                end_date = (query.get("end_date") or [default_form_values()["end_date"]])[0]
                view = (query.get("view") or ["All"])[0]
                self._send_json(build_holiday_payload(start_date, end_date, view=view, user_id=current_user.id))
                return
            if route == "/api/posting_dates":
                current_user = self._require_user()
                query = parse_qs(parsed.query)
                start_date = (query.get("start_date") or [default_form_values()["start_date"]])[0]
                end_date = (query.get("end_date") or [default_form_values()["end_date"]])[0]
                self._send_json(build_posting_date_payload(start_date, end_date, user_id=current_user.id))
                return
            self.send_error(HTTPStatus.NOT_FOUND, "Not found")
        except Exception as error:
            self._send_error(error)

    def do_POST(self) -> None:  # noqa: N802
        try:
            route = urlparse(self.path).path

            if route == "/api/login":
                payload = self._read_json_body()
                username = str(payload.get("username", ""))
                client_ip = self._client_ip()
                assert_login_not_rate_limited(client_ip, username)
                try:
                    session = auth_store().login(
                        username,
                        str(payload.get("password", "")),
                        str(payload.get("device_id", "")),
                        str(payload.get("device_label", "")),
                    )
                except ValueError:
                    # Only invalid-credential failures count toward the lockout,
                    # mirroring the PHP app. Device-limit rejections raise
                    # PermissionError and are intentionally not recorded here.
                    record_login_attempt(client_ip, username, False)
                    raise
                record_login_attempt(client_ip, username, True)
                self._send_json(
                    {
                        "token": session["token"],
                        "user": session["user"].as_payload(),
                    }
                )
                auth_store().log_activity(session["user"].id, session["user"].username, "login", "Web login completed.")
                return

            if route == "/api/logout":
                current_user = self._require_user()
                token = extract_bearer_token(self.headers)
                if token:
                    auth_store().logout(token)
                auth_store().log_activity(current_user.id, current_user.username, "logout", "Web logout completed.")
                self._send_json({"ok": True})
                return

            if route == "/api/profile":
                current_user = self._require_user()
                payload = self._read_json_body()
                profile = payload.get("profile", payload)
                if not isinstance(profile, dict):
                    raise ValueError("Profile data must be a JSON object.")
                self._send_json(auth_store().save_profile(current_user, profile))
                return

            if route == "/api/profile_formats":
                current_user = self._require_user()
                payload = self._read_json_body()
                action = str(payload.get("action", "save")).strip().lower() or "save"
                if action in {"delete", "remove"} or (action == "save" and "profile" not in payload and str(payload.get("name", "")).strip()):
                    password = str(payload.get("password", "")).strip()
                    if not password or not auth_store().verify_user_password(current_user.id, password):
                        raise PermissionError("Enter your account password to delete a saved profile format.")
                    deleted = auth_store().delete_profile_format(current_user, str(payload.get("name", "")))
                    auth_store().log_activity(current_user.id, current_user.username, "delete_profile_format", f"Deleted profile format {deleted.get('name', '')}.")
                    self._send_json(deleted)
                    return
                if action != "save":
                    raise ValueError("Unknown profile format action.")

                profile = payload.get("profile")
                if not isinstance(profile, dict):
                    raise ValueError("Profile format data must be a JSON object.")
                original_name = str(payload.get("original_name", "")).strip()
                if original_name:
                    password = str(payload.get("password", "")).strip()
                    if not password or not auth_store().verify_user_password(current_user.id, password):
                        raise PermissionError("Enter your account password to edit a saved profile format.")
                saved = auth_store().save_profile_format(
                    current_user,
                    str(payload.get("name", "")),
                    profile,
                    original_name=original_name or None,
                )
                auth_store().log_activity(current_user.id, current_user.username, "save_profile_format", f"Saved profile format {saved.get('name', '')}.")
                self._send_json(saved)
                return

            if route == "/api/templates":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                template_dir = str(payload.get("template_dir", default_form_values()["template_dir"]))
                self._send_json(serialize_catalog(template_dir, current_user.id))
                return

            if route == "/api/template_upload":
                current_user = self._require_statement_user()
                fields, files = self._read_multipart_form()
                uploaded = files.get("template_file")
                if uploaded is None:
                    raise ValueError("Select a template file first.")
                saved = save_uploaded_template(
                    fields.get("kind", ""),
                    uploaded["filename"],
                    uploaded["content"],
                    fields.get("name", ""),
                    current_user.id,
                    fields.get("profile_json", "{}"),
                )
                self._send_json(
                    {
                        "saved_template": saved,
                        "templates": serialize_catalog(fields.get("template_dir", default_form_values()["template_dir"]), current_user.id),
                    }
                )
                auth_store().log_activity(current_user.id, current_user.username, "template_upload", f"{fields.get('kind', '')}:{saved.get('name', '')}")
                return

            if route == "/api/template_detail":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                detail = template_detail(
                    str(payload.get("kind", "")),
                    str(payload.get("name", "")),
                    str(payload.get("template_dir", default_form_values()["template_dir"])),
                    current_user.id,
                )
                self._send_json(detail)
                return

            if route == "/api/letterhead":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                value = save_letterhead(APP_ROOT / 'letterheads', current_user.id, str(payload.get('kind', '')), str(payload.get('name', '')), payload)
                self._send_json({'letterhead': value})
                return
            if route == "/api/template_update_batch":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                result = update_template_batch(str(payload.get('kind', '')), str(payload.get('name', '')), str(payload.get('template_dir', '')), current_user.id, payload.get('edits'))
                self._send_json(result)
                return
            if route == "/api/template_update":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                updated = update_template_item(
                    str(payload.get("kind", "")),
                    str(payload.get("name", "")),
                    str(payload.get("template_dir", default_form_values()["template_dir"])),
                    current_user.id,
                    str(payload.get("key", "")),
                    str(payload.get("text", "")),
                    payload.get("style", {}),
                )
                auth_store().log_activity(current_user.id, current_user.username, "template_update", f"{payload.get('kind', '')}:{payload.get('name', '')}:{payload.get('key', '')}")
                self._send_json(updated)
                return

            if route == "/api/template_profile_update":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                saved_profile = save_template_profile(
                    str(payload.get("kind", "")),
                    str(payload.get("name", "")),
                    str(payload.get("template_dir", default_form_values()["template_dir"])),
                    current_user.id,
                    payload.get("profile", {}),
                )
                auth_store().log_activity(current_user.id, current_user.username, "template_profile_update", f"{payload.get('kind', '')}:{payload.get('name', '')}")
                self._send_json(saved_profile)
                return

            if route == "/api/template_create_statement":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                created = create_statement_template_from_definition(
                    str(payload.get("name", "")),
                    payload.get("definition", {}),
                    payload.get("profile", {}),
                    current_user.id,
                )
                auth_store().log_activity(current_user.id, current_user.username, "template_create_statement", f"Created statement format {created.get('name', '')}.")
                self._send_json(
                    {
                        "saved_template": created,
                        "templates": serialize_catalog(str(payload.get("template_dir", default_form_values()["template_dir"])), current_user.id),
                    }
                )
                return

            if route == "/api/template_delete":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                deleted_source = delete_uploaded_template(
                    str(payload.get("kind", "")),
                    str(payload.get("name", "")),
                    current_user.id,
                    str(payload.get("template_dir", default_form_values()["template_dir"])),
                )
                self._send_json(
                    {
                        "deleted": True,
                        "source": deleted_source,
                        "templates": serialize_catalog(str(payload.get("template_dir", default_form_values()["template_dir"])), current_user.id),
                    }
                )
                auth_store().log_activity(current_user.id, current_user.username, "template_delete", f"{payload.get('kind', '')}:{payload.get('name', '')}")
                return

            if route == "/api/holidays":
                payload = self._read_json_body()
                current_user = self._require_statement_user()
                result = apply_rule_action(payload, current_user)
                auth_store().log_activity(current_user.id, current_user.username, "holiday_change", "Updated manual holidays.")
                self._send_json(result)
                return
            if route == "/api/holidays_sync":
                self._require_statement_user()
                self._read_json_body()
                self._send_json({"error": "Automatic holiday updates are disabled. Add holidays manually."}, status=HTTPStatus.GONE)
                return
            if route == "/api/posting_dates":
                payload = self._read_json_body()
                current_user = self._require_statement_user()
                result = apply_posting_date_action(payload, current_user)
                auth_store().log_activity(current_user.id, current_user.username, "posting_date_change", "Updated interest and tax dates.")
                self._send_json(result)
                return
            if route == "/api/posting_dates_sync":
                payload = self._read_json_body()
                current_user = self._require_statement_user()
                result = sync_hamropatro_posting_dates(payload, current_user)
                auth_store().log_activity(current_user.id, current_user.username, "posting_date_sync", "Synced interest and tax dates from Hamro Patro.")
                self._send_json(result)
                return

            if route == "/api/state_restore":
                payload = self._read_json_body()
                current_user = self._require_admin()
                result = restore_user_date_rules(payload, current_user)
                auth_store().log_activity(current_user.id, current_user.username, "restore_user_date_rules", f"Restored {result.get('part', '')} for user #{result.get('user_id', '')}.")
                self._send_json(result)
                return

            if route == "/api/generate":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                config = build_config(payload, user_id=current_user.id)
                result = generate_statement(config)
                response = serialize_result(config, result)
                saved = auth_store().save_statement(current_user, response, payload, None, "generated", True)
                response["statement_id"] = saved["statement_id"]
                response["source_type"] = "generated"
                response["current_user"] = auth_store().require_user(extract_bearer_token(self.headers)).as_payload()
                response["holidays"] = build_holiday_payload(config.start_date.isoformat(), config.end_date.isoformat(), view="All", user_id=current_user.id)
                response["history"] = auth_store().statement_history(current_user)
                auth_store().log_activity(current_user.id, current_user.username, "generate_statement", f"Generated statement with {len(response['rows'])} rows.")
                self._send_json(response)
                return

            if route == "/api/import_statement":
                current_user = self._require_statement_user()
                fields, files = self._read_multipart_form()
                uploaded = files.get("statement_file")
                if uploaded is None:
                    raise ValueError("Select an Excel statement file first.")
                config_payload_raw = fields.get("config", "{}")
                try:
                    config_payload = json.loads(config_payload_raw)
                except json.JSONDecodeError as error:
                    raise ValueError("Import request config data is invalid.") from error
                if not isinstance(config_payload, dict):
                    raise ValueError("Import request config data is invalid.")

                temp_dir = Path(tempfile.mkdtemp(prefix="web_statement_import_", dir=str(WEB_TEMP_ROOT)))
                try:
                    upload_path = temp_dir / Path(uploaded["filename"]).name
                    upload_path.write_bytes(uploaded["content"])
                    imported = import_xlsx_statement(upload_path, config_payload)
                finally:
                    shutil.rmtree(temp_dir, ignore_errors=True)

                form_payload = imported["form_values"]
                config = build_config(form_payload, user_id=current_user.id)
                imported_rows = imported["edited_rows"]
                validation_errors = validate_edited_statement(config, imported_rows)
                result = recalculate_edited_statement(config, imported_rows)
                response = serialize_result(config, result)
                if not config.prepend_statement_mode:
                    response["rows"] = serialize_preview_rows(imported_rows)
                    response["summary"]["row_count"] = len(imported_rows)
                response["validation_errors"] = validation_errors
                saved = auth_store().save_statement(current_user, response, form_payload, None, "imported", current_user.access_mode != "check_only")
                response["statement_id"] = saved["statement_id"]
                response["source_type"] = "imported"
                response["current_user"] = auth_store().require_user(extract_bearer_token(self.headers)).as_payload()
                response["import_report"] = imported["report"]
                response["config"] = form_payload
                response["holidays"] = build_holiday_payload(config.start_date.isoformat(), config.end_date.isoformat(), view="All", user_id=current_user.id)
                response["history"] = auth_store().statement_history(current_user)
                auth_store().log_activity(current_user.id, current_user.username, "import_statement", f"Imported statement with {len(response['rows'])} rows.")
                self._send_json(response)
                return

            if route == "/api/recalculate_statement":
                payload = self._read_json_body()
                current_user = self._require_statement_user()
                form_payload = payload.get("config")
                edited_rows = payload.get("edited_rows")
                if not isinstance(form_payload, dict):
                    raise ValueError("Statement edit request is missing config data.")
                if not isinstance(edited_rows, list):
                    raise ValueError("Statement edit request is missing editable rows.")
                if edited_rows:
                    last_row = edited_rows[-1] if isinstance(edited_rows[-1], dict) else {}
                    if float(last_row.get("debit", 0) or 0) > 0 or float(last_row.get("credit", 0) or 0) > 0:
                        form_payload["closing_row_mode"] = "transaction_allowed"
                config = build_config(form_payload, user_id=current_user.id)
                result = recalculate_edited_statement(config, edited_rows)
                response = serialize_result(config, result)
                statement_id = int(payload.get("statement_id", 0) or 0) or None
                source_type = str(payload.get("source_type", "edited")).strip() or "edited"
                saved = auth_store().save_statement(current_user, response, form_payload, statement_id, source_type, False)
                response["statement_id"] = saved["statement_id"]
                response["source_type"] = source_type
                response["current_user"] = auth_store().require_user(extract_bearer_token(self.headers)).as_payload()
                response["holidays"] = build_holiday_payload(config.start_date.isoformat(), config.end_date.isoformat(), view="All", user_id=current_user.id)
                response["history"] = auth_store().statement_history(current_user)
                auth_store().log_activity(current_user.id, current_user.username, "recalculate_statement", "Updated statement after edit.")
                self._send_json(response)
                return

            if route == "/api/delete_statement":
                current_user = self._require_user()
                payload = self._read_json_body()
                deleted = auth_store().delete_statement(current_user, int(payload.get("statement_id", 0) or 0))
                deleted["history"] = auth_store().statement_history(current_user)
                auth_store().log_activity(current_user.id, current_user.username, "delete_statement", f"Deleted saved statement #{int(payload.get('statement_id', 0) or 0)}.")
                self._send_json(deleted)
                return

            if route == "/api/validate_statement":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                form_payload = payload.get("config")
                edited_rows = payload.get("edited_rows")
                if not isinstance(form_payload, dict):
                    raise ValueError("Statement check request is missing config data.")
                if not isinstance(edited_rows, list):
                    raise ValueError("Statement check request is missing editable rows.")
                if edited_rows:
                    last_row = edited_rows[-1] if isinstance(edited_rows[-1], dict) else {}
                    if float(last_row.get("debit", 0) or 0) > 0 or float(last_row.get("credit", 0) or 0) > 0:
                        form_payload["closing_row_mode"] = "transaction_allowed"
                config = build_config(form_payload, user_id=current_user.id)
                errors = validate_transaction_dates(config, edited_rows) if payload.get("dates_only") is True else validate_edited_statement(config, edited_rows)
                self._send_json(
                    {
                        "ok": not errors,
                        "errors": errors,
                        "checked_rows": len(edited_rows),
                        "config": form_payload,
                    }
                )
                return

            if route == "/api/users":
                current_user = self._require_admin()
                payload = self._read_json_body()
                created_user = auth_store().create_user(
                    str(payload.get("username", "")),
                    str(payload.get("password", "")),
                    str(payload.get("role", "user")),
                    str(payload.get("access_mode", "unlimited")),
                    int(str(payload.get("remaining_statements", "")).strip()) if str(payload.get("remaining_statements", "")).strip() else None,
                    str(payload.get("valid_until", "")),
                    str(payload.get("full_name", "")),
                    str(payload.get("address", "")),
                    str(payload.get("mobile_number", "")),
                    str(payload.get("email", "")),
                    str(payload.get("gender", "")),
                )
                auth_store().log_activity(current_user.id, current_user.username, "create_user", f"Created user {created_user.username}.")
                self._send_json({"created_user": created_user.as_payload(), "users": auth_store().list_users()})
                return

            if route == "/api/update_user_access":
                current_user = self._require_admin()
                payload = self._read_json_body()
                admin_password = str(payload.get("admin_password", "")).strip()
                if not admin_password or not auth_store().verify_user_password(current_user.id, admin_password):
                    raise PermissionError("Enter your admin password to update a user.")
                updated_user = auth_store().update_user_access(
                    int(payload.get("user_id", 0) or 0),
                    str(payload.get("access_mode", "unlimited")),
                    int(str(payload.get("remaining_statements", "")).strip()) if str(payload.get("remaining_statements", "")).strip() else None,
                    str(payload.get("valid_until", "")),
                )
                auth_store().log_activity(current_user.id, current_user.username, "update_user_access", f"Updated access for {updated_user.username}.")
                self._send_json({"updated_user": updated_user.as_payload(), "users": auth_store().list_users()})
                return

            if route == "/api/change_user_password":
                current_user = self._require_admin()
                payload = self._read_json_body()
                admin_password = str(payload.get("admin_password", "")).strip()
                if not admin_password or not auth_store().verify_user_password(current_user.id, admin_password):
                    raise PermissionError("Enter your admin password to change a user password.")
                updated_user = auth_store().update_user_password(
                    int(payload.get("user_id", 0) or 0),
                    str(payload.get("new_password", "")),
                )
                auth_store().log_activity(current_user.id, current_user.username, "change_user_password", f"Changed password for {updated_user.username}.")
                self._send_json({"updated_user": updated_user.as_payload(), "users": auth_store().list_users()})
                return

            if route == "/api/delete_user":
                current_user = self._require_admin()
                payload = self._read_json_body()
                admin_password = str(payload.get("admin_password", "")).strip()
                if not admin_password or not auth_store().verify_user_password(current_user.id, admin_password):
                    raise PermissionError("Enter your admin password to delete a user.")
                deleted_user = auth_store().delete_user(
                    int(payload.get("user_id", 0) or 0),
                    current_user.id,
                )
                auth_store().log_activity(current_user.id, current_user.username, "delete_user", f"Deleted user {deleted_user.username}.")
                self._send_json(
                    {
                        "deleted_user": deleted_user.as_payload(),
                        "users": auth_store().list_users(),
                        "history": auth_store().statement_history(current_user),
                    }
                )
                return

            if route == "/api/devices":
                current_user = self._require_user()
                payload = self._read_json_body()
                result = auth_store().remove_device(
                        current_user,
                        int(payload.get("user_id", 0) or current_user.id),
                        str(payload.get("device_id", "")),
                        str(payload.get("password", "")),
                    )
                auth_store().log_activity(current_user.id, current_user.username, "remove_device", "Removed a saved device serial.")
                self._send_json(result)
                return

            if route == "/api/export":
                current_user = self._require_statement_user()
                payload = self._read_json_body()
                self._handle_export(current_user, payload)
                auth_store().log_activity(current_user.id, current_user.username, f"export_{str(payload.get('export_kind', '')).strip()}", f"Exported {str(payload.get('export_kind', '')).strip()} using {str(payload.get('export_mode', '')).strip()} mode.")
                return

            if route == "/api/selftest":
                self._require_statement_user()
                self._read_json_body()
                test_result = run_tests()
                self._send_json(
                    {
                        "tests_run": test_result.testsRun,
                        "failures": len(test_result.failures),
                        "errors": len(test_result.errors),
                        "successful": test_result.wasSuccessful(),
                    }
                )
                return

            self.send_error(HTTPStatus.NOT_FOUND, "Not found")
        except Exception as error:
            self._send_error(error)

    def _read_json_body(self) -> dict[str, object]:
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length) if length else b"{}"
        if not raw.strip():
            return {}
        payload = json.loads(raw.decode("utf-8"))
        if not isinstance(payload, dict):
            raise ValueError("Request payload must be a JSON object.")
        return payload

    def _read_multipart_form(self) -> tuple[dict[str, str], dict[str, dict[str, object]]]:
        content_type = str(self.headers.get("Content-Type", "")).strip()
        if "multipart/form-data" not in content_type.lower():
            raise ValueError("This request must be sent as multipart form data.")
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length) if length else b""
        if not raw:
            return {}, {}
        envelope = (
            f"Content-Type: {content_type}\r\n"
            "MIME-Version: 1.0\r\n\r\n"
        ).encode("utf-8") + raw
        message = BytesParser(policy=email_policy).parsebytes(envelope)
        fields: dict[str, str] = {}
        files: dict[str, dict[str, object]] = {}
        for part in message.iter_parts():
            if part.get_content_disposition() != "form-data":
                continue
            field_name = str(part.get_param("name", header="content-disposition") or "").strip()
            if not field_name:
                continue
            filename = part.get_filename()
            payload_bytes = part.get_payload(decode=True) or b""
            if filename:
                files[field_name] = {"filename": filename, "content": payload_bytes}
                continue
            charset = part.get_content_charset() or "utf-8"
            fields[field_name] = payload_bytes.decode(charset, errors="replace")
        return fields, files

    def _handle_export(self, current_user: AuthUser, payload: dict[str, object]) -> None:
        export_kind = str(payload.get("export_kind", "")).strip()
        export_mode = str(payload.get("export_mode", "")).strip()
        preview_html = bool(payload.get("print_preview_html"))
        if export_kind not in {"statement", "certificate"}:
            raise ValueError("Export kind must be statement or certificate.")
        if export_mode not in {"template", "normal"}:
            raise ValueError("Export mode must be template or normal.")

        config, result, exchange_rate, export_payload = build_export_input(payload, current_user.id)
        del exchange_rate

        if export_mode == "normal":
            suffix = ".xls" if export_kind == "statement" else ".doc"
            template_name = "normal"
            body = (
                build_normal_statement_bytes(export_payload)
                if export_kind == "statement"
                else build_normal_certificate_bytes(export_payload)
            )
        else:
            template_dir = str(payload.get("template_dir") or payload.get("config", {}).get("template_dir") or default_form_values()["template_dir"])
            template_name = str(payload.get("template_name", "")).strip()
            exporter = export_statement if export_kind == "statement" else export_certificate
            selected = resolve_template_entry(export_kind, template_name, template_dir, current_user.id)
            if selected is None:
                raise ValueError(f"Selected {export_kind} template was not found.")
            suffix = selected.path.suffix

        filename = default_output_name(export_kind, config.customer_name, template_name or "normal", result.issue_date, suffix)
        if export_mode == "normal":
            self._send_download(sample_html(body.decode("utf-8")).encode("utf-8"), filename)
            return

        WEB_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
        temp_dir = WEB_TEMP_ROOT / f"web_statement_export_{uuid4().hex}"
        temp_dir.mkdir(parents=True, exist_ok=True)
        try:
            output_path = temp_dir / filename
            source_path = editable_office_copy(selected.path)
            working_output = output_path.with_suffix(source_path.suffix)
            exporter(source_path, working_output, export_payload)
            if working_output != output_path and not preview_html:
                convert_office(working_output, output_path)
            if preview_html:
                if export_kind == "statement" and working_output.suffix.lower() == ".xlsx":
                    self._send_download(build_excel_preview_html_bytes(working_output, template_name or "Statement"), Path(filename).with_suffix(".html").name)
                elif export_kind == "certificate" and working_output.suffix.lower() == ".docx":
                    self._send_download(build_docx_preview_html_bytes(working_output, template_name or "Balance Certificate"), Path(filename).with_suffix(".html").name)
                else:
                    raise ValueError('This file could not be converted into a printable document.')
            else:
                self._send_download(output_path.read_bytes(), filename)
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def _client_ip(self) -> str:
        forwarded = str(self.headers.get("CF-Connecting-IP", "")).strip()
        if forwarded:
            return forwarded
        try:
            return str(self.client_address[0])
        except Exception:
            return "unknown"

    def _security_headers(self) -> None:
        csp = "; ".join(
            [
                "default-src 'self'",
                "script-src 'self'",
                "style-src 'self' 'unsafe-inline'",
                "img-src 'self' data:",
                "font-src 'self' data:",
                "connect-src 'self'",
                "object-src 'none'",
                "base-uri 'self'",
                "form-action 'self'",
                "frame-ancestors 'none'",
            ]
        )
        self.send_header("Content-Security-Policy", csp)
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header(
            "Permissions-Policy",
            "camera=(), microphone=(), geolocation=(), payment=(), usb=(), interest-cohort=()",
        )
        self.send_header("Cross-Origin-Opener-Policy", "same-origin")
        self.send_header("Cross-Origin-Resource-Policy", "same-origin")
        self.send_header("X-Permitted-Cross-Domain-Policies", "none")
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")

    def _require_user(self) -> AuthUser:
        try:
            return auth_store().require_user(extract_bearer_token(self.headers))
        except PermissionError as error:
            raise PermissionError(str(error)) from error

    def _require_admin(self) -> AuthUser:
        try:
            return auth_store().require_admin(extract_bearer_token(self.headers))
        except PermissionError as error:
            raise PermissionError(str(error)) from error

    def _require_statement_user(self) -> AuthUser:
        try:
            return auth_store().require_statement_user(extract_bearer_token(self.headers))
        except PermissionError as error:
            raise PermissionError(str(error)) from error

    def _send_json(self, payload: dict[str, object], status: int = 200) -> None:
        body = json_bytes(payload)
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self._security_headers()
        self.end_headers()
        self.wfile.write(body)

    def _send_error(self, error: Exception) -> None:
        detail = str(error).strip() or error.__class__.__name__
        traceback.print_exc()
        if isinstance(error, PermissionError):
            status = 403
        elif isinstance(error, (ValueError, RuntimeError)):
            status = 400
        else:
            status = 500
        payload: dict[str, object] = {"error": detail}
        if _is_exchange_rate_error(detail):
            payload["requires_manual_rate"] = True
        self._send_json(payload, status=status)

    def _serve_file(self, path: Path, content_type: str | None = None, base: Path | None = None) -> None:
        if base is not None:
            # Reject any request that resolves outside the allowed base directory
            # (e.g. "/static/../auth_store.py" or the SQLite auth database).
            try:
                resolved = path.resolve()
                base_resolved = base.resolve()
            except OSError:
                self.send_error(HTTPStatus.NOT_FOUND, "File not found")
                return
            if resolved != base_resolved and base_resolved not in resolved.parents:
                self.send_error(HTTPStatus.NOT_FOUND, "File not found")
                return
            path = resolved
        if not path.exists() or not path.is_file():
            self.send_error(HTTPStatus.NOT_FOUND, "File not found")
            return
        body = path.read_bytes()
        guessed = content_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        self.send_response(200)
        self.send_header("Content-Type", guessed)
        self.send_header("Content-Length", str(len(body)))
        self._security_headers()
        self.end_headers()
        self.wfile.write(body)

    def _send_download(self, body: bytes, filename: str) -> None:
        content_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(body)))
        self._security_headers()
        self.end_headers()
        self.wfile.write(body)


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run the Web Statement Generator app.")
    parser.add_argument("--host", default="127.0.0.1", help="Host interface to bind. Default: 127.0.0.1")
    parser.add_argument("--port", type=int, default=8050, help="Port to bind. Default: 8050")
    parser.add_argument("--open", action="store_true", help="Open the app in the default web browser.")
    return parser


def main() -> None:
    args = build_argument_parser().parse_args()
    server = ThreadingHTTPServer((args.host, args.port), StatementWebHandler)
    url = f"http://{args.host}:{args.port}/"
    print(f"Web Statement Generator running at {url}")
    if args.open:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping server...")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
